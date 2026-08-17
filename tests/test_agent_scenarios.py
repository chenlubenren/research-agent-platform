from __future__ import annotations

import asyncio
import json
from pathlib import Path

from pptx import Presentation

from research_agent_platform import agent as agent_module
from research_agent_platform.agent import ResearchAgentService

from .conftest import run


def test_direct_chat_identity_reply(service: ResearchAgentService):
    result = run(service.chat(None, "你是谁"))
    assert result["status"] == "idle"
    assert result["task_id"] == ""
    assert result["checkpoint"] is None
    assert "科研" in result["text"]


def test_present_runs_without_routine_checkpoint(service: ResearchAgentService, isolated_env):
    start = run(service.chat(None, "/present 做一个中文汇报"))
    task_id = start["task_id"]

    assert start["status"] == "completed"
    assert start["command"] == "/present"
    assert start["checkpoint"] is None
    assert Path(isolated_env["state_root"], "langgraph-checkpoints.pkl").exists()

    task = service.get_task(task_id)
    assert task is not None
    assert task.status == "completed"
    assert task.user_id == "local"
    artifact_root = Path(task.artifact_root).as_posix()
    assert f"local/{task.session_id}" in artifact_root
    assert Path(task.artifact_root, "presentation", "SLIDES_OUTLINE.md").exists()
    assert Path(task.artifact_root, "bib").is_dir()
    assert Path(task.artifact_root, "code").is_dir()
    assert Path(task.artifact_root, "figures").is_dir()
    assert Path(task.artifact_root, "paper").is_dir()
    assert Path(task.artifact_root, "Content").is_dir()
    wiki_note = Path(task.artifact_root, "wiki", "agent-notes", f"{task.task_id}.md")
    assert wiki_note.exists()
    assert "Status: completed" in wiki_note.read_text(encoding="utf-8")
    assert any(artifact.relative_path == f"wiki/agent-notes/{task.task_id}.md" for artifact in task.artifacts)
    assert all(Path(note.removeprefix("Wiki note: ")).is_relative_to(Path(task.artifact_root)) for note in task.notes if note.startswith("Wiki note: "))
    assert len(task.approvals) == 0
    assert task.current_stage_name == "qa_brief"
    assert any(artifact.relative_path == "presentation/SLIDES_OUTLINE.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/SLIDE_CONTENT.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/SPEAKER_NOTES.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/QA_BRIEF.md" for artifact in task.artifacts)
    assert any(artifact.relative_path == "presentation/STAGE_REPORT.pptx" for artifact in task.artifacts)
    assert Path(task.artifact_root, "presentation", "slides", "SLIDE_01.png").exists()
    assert Path(task.artifact_root, "Content", "PRESENTATION_TEMPLATE.json").exists()
    notes_json = Path(task.artifact_root, "Content", "SPEAKER_NOTES.json")
    assert notes_json.exists()
    assert json.loads(notes_json.read_text(encoding="utf-8"))[0]["notes"]
    deck = Presentation(Path(task.artifact_root, "presentation", "STAGE_REPORT.pptx"))
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides)


def test_stage_prompt_enforces_session_workspace_boundary(
    service: ResearchAgentService, monkeypatch
):
    system_prompts: list[str] = []

    async def capture_prompt(*, system_prompt, user_prompt, model=None, temperature=0.3):
        system_prompts.append(system_prompt)
        return "# Research artifact\n\n- Session-scoped output."

    monkeypatch.setattr(agent_module, "generate_text", capture_prompt)
    session = service.store.create_session()
    result = run(service.chat(session.session_id, "/review 海绵城市"))
    completed = service.get_task(result["task_id"])

    assert completed is not None
    expected_root = str(Path(completed.artifact_root).resolve())
    assert system_prompts
    assert all(expected_root in prompt for prompt in system_prompts)
    assert all("Never write or propose writing research files" in prompt for prompt in system_prompts)
    assert all("wiki/ for session memory and task notes" in prompt for prompt in system_prompts)
    note_path = Path(completed.artifact_root, "wiki", "agent-notes", f"{completed.task_id}.md")
    assert note_path.exists()
    assert not Path(service.artifacts.root).parent.joinpath("research-wiki", "agent-notes", f"{completed.task_id}.md").exists()


def test_paper_talk_uses_paper_delivery_name(service: ResearchAgentService):
    session = service.store.create_session()
    Path(session.workspace_root, "paper").mkdir(parents=True, exist_ok=True)
    Path(session.workspace_root, "paper", "FINAL_PAPER.md").write_text(
        "# Final Paper\n\nValidated result.", encoding="utf-8"
    )
    start = run(service.chat(session.session_id, "/present 论文汇报"))
    task = service.get_task(start["task_id"])
    assert task is not None
    assert Path(task.artifact_root, "presentation", "slides").exists()

    result = start
    task = service.get_task(task.task_id)

    assert result["status"] == "completed"
    assert task is not None
    assert any(artifact.relative_path == "presentation/PAPER_TALK.pptx" for artifact in task.artifacts)
    assert Path(task.artifact_root, "Content", "PRESENTATION_SOURCE_INDEX.md").read_text(
        encoding="utf-8"
    ).find("paper/FINAL_PAPER.md") >= 0


def test_present_source_set_is_frozen_at_task_start(service: ResearchAgentService):
    session = service.store.create_session()
    Path(session.workspace_root, "paper").mkdir(parents=True, exist_ok=True)
    Path(session.workspace_root, "paper", "INITIAL.md").write_text("initial", encoding="utf-8")

    start = run(service.chat(session.session_id, "/present --source workspace 论文汇报"))
    task = service.get_task(start["task_id"])
    assert task is not None
    assert task.presentation_source is not None
    assert "paper/INITIAL.md" in task.presentation_source.source_refs

    Path(session.workspace_root, "paper", "LATE.md").write_text("late", encoding="utf-8")
    assert start["status"] == "completed"
    selection = Path(
        task.artifact_root,
        "Content",
        "PRESENTATION_SOURCE_SELECTION.json",
    ).read_text(encoding="utf-8")

    assert "paper/INITIAL.md" in selection
    assert "paper/LATE.md" not in selection


def test_present_only_checkpoints_for_complete_blocking_decision(service: ResearchAgentService, monkeypatch):
    original_generate_text = agent_module.generate_text

    async def generate_with_choices(*, system_prompt, user_prompt, model=None, temperature=0.3):
        content = await original_generate_text(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=model,
            temperature=temperature,
        )
        if "Current stage: Slides Outline" in user_prompt:
            return content + (
                "\n\n## Decision Required\n"
                "Blocking: Yes\n"
                "Question: Choose the disclosure scope.\n"
                "Why user input is necessary: This changes which unpublished result is disclosed externally.\n"
                "Option A: Include the unpublished result.\n"
                "Option B: Exclude the unpublished result.\n"
                "Recommended Default: Option B.\n"
            )
        return content

    monkeypatch.setattr(agent_module, "generate_text", generate_with_choices)
    start = run(service.chat(None, "/present 做一个中文汇报"))
    task_id = start["task_id"]

    assert start["status"] == "waiting_human"
    assert start["checkpoint"]["title"] == "Presentation Outline Approval"
    rejected = run(service.reject_task(task_id, "请选择更详细的叙事路线"))
    task = service.get_task(task_id)

    assert rejected["status"] == "waiting_human"
    assert rejected["checkpoint"]["title"] == "Presentation Outline Approval"
    assert task is not None
    assert task.status == "waiting_human"
    assert len(task.approvals) == 2
    assert task.approvals[0].status == "rejected"
    assert task.approvals[1].status == "pending"
    assert any("打回" in item for item in task.progress_log)


def test_present_does_not_checkpoint_for_unstructured_preferences(service: ResearchAgentService, monkeypatch):
    original_generate_text = agent_module.generate_text

    async def generate_with_preferences(*, system_prompt, user_prompt, model=None, temperature=0.3):
        content = await original_generate_text(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=model,
            temperature=temperature,
        )
        if "Current stage: Slides Outline" in user_prompt:
            return content + "\n\n## Decision Required\n- Option A: concise talk\n- Option B: detailed talk\n"
        return content

    monkeypatch.setattr(agent_module, "generate_text", generate_with_preferences)
    start = run(service.chat(None, "/present 做一个中文汇报"))

    assert start["status"] == "completed"
    assert start["checkpoint"] is None


def test_present_does_not_checkpoint_for_markdown_none_decision(service: ResearchAgentService, monkeypatch):
    original_generate_text = agent_module.generate_text

    async def generate_with_none(*, system_prompt, user_prompt, model=None, temperature=0.3):
        content = await original_generate_text(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=model,
            temperature=temperature,
        )
        if "Current stage: Slides Outline" in user_prompt:
            return content + "\n\n## Open Design Questions\n- Option A\n- Option B\n\n## Decision Required\n\n**None.**\n"
        return content

    monkeypatch.setattr(agent_module, "generate_text", generate_with_none)
    start = run(service.chat(None, "/present 做一份论文汇报"))

    assert start["status"] == "completed"
    assert start["checkpoint"] is None


def test_present_respects_explicit_request_to_wait_for_approval(service: ResearchAgentService):
    start = run(service.chat(None, "/present 先让我审核，批准后再继续生成"))

    assert start["status"] == "waiting_human"
    assert start["checkpoint"]["title"] == "Presentation Outline Approval"


def test_resume_after_restart(service: ResearchAgentService, isolated_env):
    start = run(service.chat(None, "/present 做一个中文汇报"))
    task_id = start["task_id"]

    fresh_service = ResearchAgentService()
    task = fresh_service.get_task(task_id)

    assert start["status"] == "completed"
    assert task is not None
    assert task.status == "completed"
    assert Path(isolated_env["state_root"], "langgraph-checkpoints.pkl").exists()
    assert Path(task.artifact_root).exists()


def test_resume_after_interruption_during_presentation_delivery(service, monkeypatch):
    async def scenario():
        original_generate_image = agent_module.generate_image
        image_started = asyncio.Event()
        release_image = asyncio.Event()

        async def delayed_generate_image(**kwargs):
            image_started.set()
            await release_image.wait()
            return await original_generate_image(**kwargs)

        monkeypatch.setattr(agent_module, "generate_image", delayed_generate_image)
        running = asyncio.create_task(service.chat(None, "/present 做一个中文汇报"))
        await image_started.wait()
        tasks = service.list_tasks()
        assert len(tasks) == 1
        task_id = tasks[0].task_id
        running.cancel()
        try:
            await running
        except asyncio.CancelledError:
            service.record_task_failure(task_id, asyncio.CancelledError())

        interrupted = service.get_task(task_id)
        assert interrupted is not None
        assert interrupted.status == "running"
        assert interrupted.current_stage_name == "presentation_delivery"

        monkeypatch.setattr(agent_module, "generate_image", original_generate_image)
        fresh_service = ResearchAgentService()
        resumed = await fresh_service.continue_task(task_id)
        completed = fresh_service.get_task(task_id)

        assert resumed["status"] == "completed"
        assert completed is not None
        assert completed.status == "completed"
        assert Path(completed.artifact_root, "presentation", "STAGE_REPORT.pptx").exists()

    asyncio.run(scenario())


def test_idea_verification_runs_without_routine_checkpoint(service: ResearchAgentService, monkeypatch):
    original_generate_text = agent_module.generate_text

    async def generate_idea(*, system_prompt, user_prompt, model=None, temperature=0.3):
        content = await original_generate_text(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=model,
            temperature=temperature,
        )
        if "Current stage: Idea Candidates" in user_prompt:
            return (
                "# Idea Candidates\n\n"
                "## Problem Frame\n- Topic.\n\n"
                "## Candidate Ideas\n- A\n- B\n- C\n\n"
                "## Comparative Assessment\n- A best.\n\n"
                "## Recommended Candidate\n- A\n\n"
                "## Risks and Unknowns\n- None.\n\n"
                "## Decision Required\nNone."
            )
        if "Current stage: Idea Verification" in user_prompt:
            return (
                "# Idea Verification\n\n"
                "## Candidate Under Review\n- A\n\n"
                "## Closest Prior Work\n- Prior.\n\n"
                "## Novelty Stress Test\n- Pass.\n\n"
                "## Feasibility Stress Test\n- Pass.\n\n"
                "## Disconfirming Evidence\n- None.\n\n"
                "## Unresolved Questions\n- None.\n\n"
                "## Verification Verdict\n- Keep."
            )
        if "Current stage: Final Idea" in user_prompt:
            return (
                "# 最终研究 Idea：示例方向\n\n"
                "## 摘要\n- Topic.\n\n"
                "## 1. 引言\n- Background.\n\n"
                "## 2. 相关工作\n- Prior.\n\n"
                "## 3. 研究问题与核心假设\n- Gap.\n\n"
                "## 4. 方法思路\n- A.\n\n"
                "## 5. 实验方案\n- Experiment.\n\n"
                "## 6. 预期贡献与可证伪预测\n- C.\n\n"
                "## 7. 局限、风险与不确定性\n- F.\n\n"
                "## 8. 结论与下一步\n- Ready.\n\n"
                "## 参考文献与证据\n- D."
            )
        return content

    monkeypatch.setattr(agent_module, "generate_text", generate_idea)
    start = run(service.chat(None, "/idea 给我一个创新方向"))

    assert start["status"] == "completed"
    assert start["checkpoint"] is None
