from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation

from research_agent_platform.presentation import (
    assemble_image_deck,
    assemble_mixed_deck,
    build_slide_prompt,
    collect_workspace_evidence,
    compose_slide_preview,
    crop_slide_image,
    enforce_presentation_page_mix,
    extract_presentation_assets,
    parse_slide_content,
    parse_speaker_notes,
    PRESENTATION_IMAGE_SIZE,
    presentation_design_spec_to_json,
    reconcile_slide_specs,
    resolve_presentation_source_config,
    resolve_slide_source_image,
    select_slide_asset,
    select_presentation_mode,
    select_presentation_template,
    PresentationAsset,
    SlideRender,
    SlideSpec,
    WorkspaceEvidence,
)
from research_agent_platform.artifacts.store import ArtifactStore
from research_agent_platform.models import UploadBatchRecord


def test_paper_mode_prioritizes_paper_outputs(tmp_path: Path):
    for name in ("paper", "plan", "figures", "Content"):
        (tmp_path / name).mkdir()
    (tmp_path / "paper" / "FINAL.md").write_text("paper evidence", encoding="utf-8")
    (tmp_path / "plan" / "EXPERIMENT_PLAN.md").write_text("plan evidence", encoding="utf-8")

    mode = select_presentation_mode("论文汇报", tmp_path)
    evidence = collect_workspace_evidence(tmp_path, mode)

    assert mode == "paper"
    assert evidence.text.index("paper/FINAL.md") < evidence.text.index("plan/EXPERIMENT_PLAN.md")
    assert select_presentation_template("论文汇报", mode).name == "paper-talk"


def test_slide_content_parsing_and_deck_assembly(tmp_path: Path):
    content = (
        "# Talk\n\n"
        "## Slide 1: Problem\n"
        "### Main Message\n- Define the gap.\n"
        "### Source Files\n- `paper/FINAL.md`\n\n"
        "## Slide 2: Evidence\n"
        "### Main Message\n- Show the result.\n"
        "### Source Files\n- `figures/result.png`\n"
    )
    slides = parse_slide_content(content)
    assert [slide.title for slide in slides] == ["Problem", "Evidence"]
    assert slides[1].source_files == ("figures/result.png",)

    paths: list[Path] = []
    for number in (1, 2):
        raw = io.BytesIO()
        Image.new("RGB", (1536, 1024), color=(240, 240, 240)).save(raw, format="PNG")
        cropped = crop_slide_image(raw.getvalue())
        path = tmp_path / f"slide-{number}.png"
        path.write_bytes(cropped)
        with Image.open(path) as image:
            assert image.size == (1536, 864)
        paths.append(path)

    deck_bytes = assemble_image_deck(paths, title="Test Talk")
    deck = Presentation(io.BytesIO(deck_bytes))
    assert len(deck.slides) == 2
    assert deck.slide_width / deck.slide_height == 16 / 9


def test_image2_fit_preserves_top_and_bottom_content(tmp_path: Path):
    source = Image.new("RGB", (1536, 1024), color=(20, 20, 20))
    pixels = source.load()
    for x in range(source.width):
        pixels[x, 0] = (220, 40, 40)
        pixels[x, source.height - 1] = (40, 80, 220)
    raw = io.BytesIO()
    source.save(raw, format="PNG")

    fitted = Image.open(io.BytesIO(crop_slide_image(raw.getvalue())))
    assert fitted.size == (1536, 864)
    top_pixel = fitted.getpixel((768, 0))
    bottom_pixel = fitted.getpixel((768, 863))
    assert top_pixel[0] > 150 and top_pixel[0] > top_pixel[2]
    assert bottom_pixel[2] > 150 and bottom_pixel[2] > bottom_pixel[0]


def test_presentation_image_generation_requests_native_widescreen_size():
    assert PRESENTATION_IMAGE_SIZE == "1536x864"


def test_image2_prompt_uses_user_topic_not_stage_report_as_subject():
    slide = SlideSpec(
        number=1,
        title="可解释机器学习的研究问题",
        content="### On-Slide Text\n- 解释模型预测依据。",
        source_files=(),
    )
    prompt = build_slide_prompt(
        slide,
        select_presentation_template("阶段汇报", "stage"),
        "stage",
        objective="请制作关于可解释机器学习的阶段性汇报 PPT",
    )
    assert "可解释机器学习" in prompt
    assert "只是交付格式，绝不是研究主题" in prompt
    assert "不要生成关于如何做汇报" in prompt


def test_slide_content_parsing_accepts_top_level_slide_headings():
    slides = parse_slide_content(
        "# Slide 1: Title\n**Main Message**\nFirst page.\n\n"
        "# Slide 2: Result\n**Main Message**\nSecond page.\n"
    )

    assert [slide.number for slide in slides] == [1, 2]
    assert [slide.title for slide in slides] == ["Title", "Result"]


def test_speaker_notes_are_parsed_and_embedded_in_native_ppt_notes(tmp_path: Path):
    slide = parse_slide_content(
        "## Slide 1: Result\n"
        "### Main Message\nThe model improves accuracy.\n"
        "### Render Mode\nimage2_full\n"
    )[0]
    notes_markdown = (
        "## Per-Slide Notes\n"
        "### Slide 1 — Result\n"
        "- Explain the main comparison.\n"
        "- Emphasize the validated result.\n\n"
        "## Timing\n"
        "### Slide 1 — Result\n"
        "- Suggested Duration: 1:00\n"
        "- Goal: Establish the result\n\n"
        "## Transitions\n"
        "### Slide 1 → Slide 2\n"
        "- Transition Line: Next, inspect the original evidence.\n"
    )
    notes = parse_speaker_notes(notes_markdown, [slide])
    assert "Explain the main comparison." in notes[1]
    assert "1:00" in notes[1]
    assert "Next, inspect the original evidence." in notes[1]

    generated = tmp_path / "generated.png"
    preview = tmp_path / "preview.png"
    Image.new("RGB", (1536, 864), color=(12, 34, 56)).save(generated)
    preview.write_bytes(generated.read_bytes())
    deck = Presentation(
        io.BytesIO(
            assemble_mixed_deck(
                [SlideRender(slide, generated, preview)],
                speaker_notes=notes,
            )
        )
    )

    assert deck.slides[0].notes_slide.notes_text_frame.text == notes[1]
    assert len(deck.slides[0].shapes) == 1


def test_speaker_notes_fall_back_to_slide_main_message():
    slide = parse_slide_content(
        "## Slide 1: Result\n"
        "### Main Message\nThe model improves accuracy.\n"
        "### Render Mode\nimage2_full\n"
    )[0]

    notes = parse_speaker_notes("", [slide])

    assert notes[1] == "讲解词\n本页重点：The model improves accuracy."


def test_outline_render_contract_controls_final_page_mix():
    outline = parse_slide_content(
        "## Slide 1: Cover\n"
        "### Page Type\ncover\n"
        "### Render Mode\nimage2_full\n\n"
        "## Slide 2: Result\n"
        "### Page Type\nevidence\n"
        "### Render Mode\nevidence\n"
        "### Asset IDs\n- asset_0001\n"
    )
    content = parse_slide_content(
        "## Slide 1: Cover\n"
        "### Page Type\nevidence\n"
        "### Render Mode\nevidence\n"
        "### Asset IDs\n- asset_0002\n\n"
        "## Slide 2: Result\n"
        "### Page Type\nnarrative\n"
        "### Render Mode\nimage2_full\n"
    )

    slides = reconcile_slide_specs(content, outline)

    assert [slide.render_mode for slide in slides] == ["image2_full", "evidence"]
    assert slides[0].asset_ids == ()
    assert slides[1].asset_ids == ("asset_0001",)
    design_spec = json.loads(presentation_design_spec_to_json(slides))
    assert design_spec["page_mix"] == {"total": 2, "image2_full": 1, "evidence": 1}


def test_outline_parser_ignores_numbered_sections_and_evidence_lists():
    slides = parse_slide_content(
        "# Presentation Plan\n\n"
        "### 1) Talk type and framing\n- Mode: paper\n\n"
        "## Slide 1: Cover\n- **Render Mode**: image2_full\n\n"
        "## Key Evidence per Slide\n### 1\n- No original asset\n"
    )
    assert [(slide.number, slide.title) for slide in slides] == [(1, "Cover")]


def test_page_mix_enforces_narrative_majority_without_mixing_assets():
    slides = parse_slide_content(
        "\n".join(
            f"## Slide {number}: {'Study Area' if number == 1 else 'Evidence'}\n"
            "- **Page Type**: evidence\n"
            "- **Render Mode**: evidence\n"
            f"- **Asset IDs**: asset_{number:04d}\n"
            for number in range(1, 11)
        )
    )
    normalized = enforce_presentation_page_mix(slides)
    assert sum(slide.render_mode == "image2_full" for slide in normalized) == 6
    assert sum(slide.render_mode == "evidence" for slide in normalized) == 4
    assert all(
        not slide.asset_ids
        for slide in normalized
        if slide.render_mode == "image2_full"
    )


def test_artifact_store_enforces_standard_workspace_directories(tmp_path: Path):
    store = ArtifactStore(str(tmp_path / "workspace"))
    try:
        store.task_root("task")
    except ValueError:
        pass
    else:
        raise AssertionError("Task-specific workspace roots must be rejected")
    root = store.session_root(user_id="user", session_id="session")

    store.write_text(
        "task",
        "Content/context.md",
        "context",
        kind="note",
        description="test",
        task_root=root,
    )

    for invalid in ("misc/file.md", "../outside.md"):
        try:
            store.write_text(
                "task",
                invalid,
                "invalid",
                kind="note",
                description="test",
                task_root=root,
            )
        except ValueError:
            continue
        raise AssertionError(f"Expected invalid workspace path to be rejected: {invalid}")


def test_slide_source_image_cannot_escape_workspace(tmp_path: Path):
    (tmp_path / "figures").mkdir()
    outside = tmp_path.parent / "outside.png"
    Image.new("RGB", (10, 10)).save(outside)
    slide = parse_slide_content(
        "## Slide 1: Evidence\n### Source Files\n- `figures/../../outside.png`\n"
    )[0]
    evidence = WorkspaceEvidence(mode="stage", text="", files=(), images=())

    assert resolve_slide_source_image(slide, evidence, tmp_path) is None


def test_auto_source_scope_prefers_latest_upload_batch(tmp_path: Path):
    for name in ("paper", "figures", "Content"):
        (tmp_path / name).mkdir()
    (tmp_path / "paper" / "old.pdf").write_bytes(b"old")
    (tmp_path / "paper" / "new.docx").write_bytes(b"new")
    batches = [
        UploadBatchRecord(upload_batch_id="upload_old", relative_paths=["paper/old.pdf"]),
        UploadBatchRecord(upload_batch_id="upload_new", relative_paths=["paper/new.docx"]),
    ]

    config = resolve_presentation_source_config("论文汇报", tmp_path, batches)

    assert config.resolved_scope == "attachments"
    assert config.upload_batch_id == "upload_new"
    assert config.source_refs == ["paper/new.docx"]


def test_generic_present_with_uploaded_pdf_detects_paper_mode(tmp_path: Path):
    for name in ("paper", "figures", "Content"):
        (tmp_path / name).mkdir()
    uploads = tmp_path / "paper" / "uploads"
    uploads.mkdir()
    (uploads / "paper.pdf").write_bytes(b"pdf")

    config = resolve_presentation_source_config("做一个PPT", tmp_path)

    assert config.presentation_type == "paper"
    assert config.resolved_scope == "attachments"
    assert config.source_refs == ["paper/uploads/paper.pdf"]


def test_explicit_source_path_overrides_latest_upload(tmp_path: Path):
    for name in ("paper", "figures", "Content"):
        (tmp_path / name).mkdir()
    (tmp_path / "paper" / "paper.pdf").write_bytes(b"paper")
    (tmp_path / "figures" / "result.png").write_bytes(b"result")
    batches = [UploadBatchRecord(relative_paths=["paper/paper.pdf"])]

    config = resolve_presentation_source_config(
        "/present --source selected `figures/result.png` 做汇报",
        tmp_path,
        batches,
    )

    assert config.resolved_scope == "selected"
    assert config.source_refs == ["figures/result.png"]


def test_selected_evidence_does_not_scan_other_workspace_files(tmp_path: Path):
    for name in ("paper", "plan", "figures", "Content"):
        (tmp_path / name).mkdir()
    selected = tmp_path / "paper" / "selected.md"
    ignored = tmp_path / "plan" / "ignored.md"
    selected.write_text("selected evidence", encoding="utf-8")
    ignored.write_text("ignored evidence", encoding="utf-8")

    evidence = collect_workspace_evidence(
        tmp_path,
        "paper",
        source_refs=["paper/selected.md"],
    )

    assert evidence.files == ("paper/selected.md",)
    assert "selected evidence" in evidence.text
    assert "ignored evidence" not in evidence.text


def test_csv_asset_becomes_editable_ppt_table(tmp_path: Path):
    for name in ("figures", "presentation", "Content"):
        (tmp_path / name).mkdir()
    table_path = tmp_path / "figures" / "results.csv"
    table_path.write_text("Model,R2\nRF,0.85\nSVM,0.72\n", encoding="utf-8")
    assets = extract_presentation_assets(tmp_path, ["figures/results.csv"], "figures/extracted/test")
    assert len(assets) == 1
    assert assets[0].kind == "table"
    assert assets[0].table_data[1] == ("RF", "0.85")

    base = tmp_path / "presentation" / "base.png"
    Image.new("RGB", (1536, 864), color=(250, 250, 250)).save(base)
    preview = tmp_path / "presentation" / "preview.png"
    slide = parse_slide_content(
        "## Slide 1: Results\n### On-Slide Text\n- Compare models\n"
        "### Source Files\n- `figures/results.csv`\n"
    )[0]
    compose_slide_preview(base, slide, preview, tmp_path, assets[0])
    deck = Presentation(
        io.BytesIO(assemble_mixed_deck([SlideRender(slide, base, preview, assets[0])]))
    )

    assert any(shape.has_table for shape in deck.slides[0].shapes)
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert table.cell(1, 1).text == "0.85"


def test_original_image_is_embedded_without_image2_redraw(tmp_path: Path):
    for name in ("figures", "presentation", "Content"):
        (tmp_path / name).mkdir()
    source = tmp_path / "figures" / "result.png"
    source_buffer = io.BytesIO()
    Image.new("RGB", (640, 360), color=(12, 34, 56)).save(source_buffer, format="PNG")
    source.write_bytes(source_buffer.getvalue())
    asset = PresentationAsset(
        asset_id="asset_0001",
        kind="image",
        source_path="figures/result.png",
        extracted_path="figures/result.png",
        caption="Result",
    )
    slide = parse_slide_content(
        "## Slide 1: Result\n### On-Slide Text\n- Exact evidence\n"
        "### Source Files\n- `figures/result.png`\n### Asset IDs\n- asset_0001\n"
    )[0]
    assert select_slide_asset(slide, [asset]) == asset
    prompt = build_slide_prompt(slide, select_presentation_template("阶段汇报", "stage"), "stage", asset)
    assert "不要重绘实验数据" in prompt
    assert "不要生成任何文字" in prompt

    base = tmp_path / "presentation" / "base.png"
    preview = tmp_path / "presentation" / "preview.png"
    Image.new("RGB", (1536, 864), color=(250, 250, 250)).save(base)
    compose_slide_preview(base, slide, preview, tmp_path, asset)
    deck_bytes = assemble_mixed_deck([SlideRender(slide, base, preview, asset)])
    deck_path = tmp_path / "presentation" / "deck.pptx"
    deck_path.write_bytes(deck_bytes)

    with zipfile.ZipFile(deck_path) as archive:
        media = [archive.read(name) for name in archive.namelist() if name.startswith("ppt/media/")]
    assert source.read_bytes() in media


def test_image2_page_is_standalone_and_never_receives_asset_overlay(tmp_path: Path):
    for name in ("presentation", "Content"):
        (tmp_path / name).mkdir()
    generated = tmp_path / "presentation" / "generated.png"
    preview = tmp_path / "presentation" / "preview.png"
    Image.new("RGB", (1536, 864), color=(12, 34, 56)).save(generated)
    slide = parse_slide_content(
        "## Slide 1: Contribution\n"
        "### Page Type\ncover\n"
        "### Render Mode\nimage2_full\n"
        "### On-Slide Text\n- One clear contribution\n"
    )[0]
    assert slide.render_mode == "image2_full"
    assert select_slide_asset(slide, []) is None
    compose_slide_preview(generated, slide, preview, tmp_path)
    deck = Presentation(io.BytesIO(assemble_mixed_deck([SlideRender(slide, generated, preview)])))
    assert len(deck.slides) == 1
    assert sum(shape.shape_type == 13 for shape in deck.slides[0].shapes) == 1
    assert not any(shape.has_text_frame for shape in deck.slides[0].shapes)


def test_explicit_evidence_page_is_separate_from_image2(tmp_path: Path):
    for name in ("figures", "presentation", "Content"):
        (tmp_path / name).mkdir()
    source = tmp_path / "figures" / "result.png"
    Image.new("RGB", (640, 360), color=(34, 56, 78)).save(source)
    asset = PresentationAsset(
        asset_id="asset_0001",
        kind="image",
        source_path="figures/result.png",
        extracted_path="figures/result.png",
        caption="Result",
    )
    slide = parse_slide_content(
        "## Slide 2: Result\n"
        "### Page Type\nevidence\n"
        "### Render Mode\nevidence\n"
        "### Asset IDs\n- asset_0001\n"
        "### Source Files\n- `figures/result.png`\n"
    )[0]
    assert slide.render_mode == "evidence"
    assert "不要调用 Image-2" in build_slide_prompt(
        slide, select_presentation_template("阶段汇报", "stage"), "stage", asset
    )
    preview = tmp_path / "presentation" / "preview.png"
    compose_slide_preview(None, slide, preview, tmp_path, asset)
    deck = Presentation(io.BytesIO(assemble_mixed_deck([SlideRender(slide, None, preview, asset)])))
    assert any(shape.shape_type == 13 for shape in deck.slides[0].shapes)
