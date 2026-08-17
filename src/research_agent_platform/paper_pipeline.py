from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Iterable

from .models import UploadBatchRecord, WriteSourceConfig


TEXT_EXTENSIONS = {
    ".bib",
    ".cfg",
    ".csv",
    ".enw",
    ".html",
    ".ini",
    ".ipynb",
    ".json",
    ".log",
    ".md",
    ".nbib",
    ".out",
    ".py",
    ".r",
    ".rst",
    ".ris",
    ".sh",
    ".sql",
    ".tex",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
IMAGE_EXTENSIONS = {".gif", ".jpeg", ".jpg", ".png", ".svg", ".tif", ".tiff", ".webp"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".pptx", ".xlsx"} | IMAGE_EXTENSIONS
GENERATED_PAPER_FILES = {
    "PAPER_EVIDENCE_MAP.json",
    "PAPER_SELF_REVIEW.md",
    "PAPER_DELIVERY_REPORT.json",
    "CITATION_AUDIT.json",
}
IGNORED_CONTEXT_FILES = {
    "CLOUD_SYNC.json",
    "MANIFEST.md",
    "PAPER_SOURCE_SELECTION.json",
    "PAPER_EVIDENCE_METADATA.json",
    "PRESENTATION_SOURCE_SELECTION.json",
    "PRESENTATION_SOURCE_INDEX.md",
}
WORKSPACE_ORDER = ("paper", "figures", "plan", "idea", "bib", "Content", "logs", "code", "wiki")


def resolve_write_source_config(
    objective: str,
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
    *,
    source_limit: int = 50,
) -> WriteSourceConfig:
    requested_scope = _requested_source_scope(objective)
    explicit_refs = _extract_explicit_refs(objective, workspace_root)
    batches = list(upload_batches)
    latest_batch = batches[-1] if batches else None
    latest_refs = _valid_refs(
        workspace_root,
        latest_batch.relative_paths if latest_batch else (),
    )[:source_limit]
    workspace_refs = select_write_workspace_sources(workspace_root, limit=source_limit)

    if explicit_refs:
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="selected",
            source_refs=_expand_refs(workspace_root, explicit_refs, source_limit),
            selection_reason="用户在 /write 指令中明确指定了写作材料。",
        )
    if requested_scope == "attachments":
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="attachments",
            source_refs=latest_refs,
            upload_batch_id=latest_batch.upload_batch_id if latest_batch else "",
            selection_reason="用户要求仅使用最近一次上传批次。",
        )
    if requested_scope == "session":
        session_refs = _deduplicate(
            ref
            for batch in batches
            for ref in _valid_refs(workspace_root, batch.relative_paths)
        )[:source_limit]
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="session",
            source_refs=session_refs,
            upload_batch_id=latest_batch.upload_batch_id if latest_batch else "",
            selection_reason="用户要求使用当前会话上传的全部写作材料。",
        )
    if requested_scope == "workspace":
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="workspace",
            source_refs=workspace_refs,
            selection_reason="用户要求检索工作区内的论文、实验、图表和研究上下文。",
        )
    if requested_scope == "selected":
        return WriteSourceConfig(
            requested_scope=requested_scope,
            resolved_scope="selected",
            source_refs=[],
            selection_reason="用户要求使用指定材料，但指令中没有解析到有效路径。",
        )

    combined = _deduplicate([*latest_refs, *workspace_refs])[:source_limit]
    if latest_refs:
        reason = "自动模式优先纳入最近上传材料，并补充工作区内相关实验、图表、计划和文献证据。"
        scope = "session"
    else:
        reason = "未发现本次上传或显式路径，自动冻结工作区内与论文写作相关的材料。"
        scope = "workspace"
    return WriteSourceConfig(
        requested_scope="auto",
        resolved_scope=scope,
        source_refs=combined,
        upload_batch_id=latest_batch.upload_batch_id if latest_batch and latest_refs else "",
        selection_reason=reason,
    )


def select_write_workspace_sources(workspace_root: Path, *, limit: int = 50) -> list[str]:
    ranked: list[tuple[int, int, str]] = []
    for directory_rank, directory_name in enumerate(WORKSPACE_ORDER):
        directory = workspace_root / directory_name
        if not directory.exists():
            continue
        for path in directory.rglob("*"):
            if not _is_source_file(path):
                continue
            relative = path.relative_to(workspace_root).as_posix()
            priority = directory_rank * 100
            lowered = path.stem.lower()
            if directory_name == "paper" and path.parent.name == "uploads":
                priority -= 70
            if directory_name == "paper" and any(
                term in lowered for term in ("final", "revised", "draft", "终稿", "定稿", "修订稿", "草稿")
            ):
                priority -= 50
            if directory_name == "figures" and path.suffix.lower() in IMAGE_EXTENSIONS:
                priority -= 30
            ranked.append((priority, -path.stat().st_mtime_ns, relative))
    return [relative for _, _, relative in sorted(ranked)[:limit]]


def collect_paper_evidence(
    workspace_root: Path,
    source_refs: Iterable[str],
    *,
    total_limit: int = 40000,
    query: str = "",
    prioritize_research_sections: bool = False,
) -> list[dict]:
    records: list[dict] = []
    used = 0
    for relative in source_refs:
        source = workspace_root / relative
        if not _is_source_file(source):
            continue
        suffix = source.suffix.lower()
        extracted = _extract_source_blocks(source)
        if suffix == ".pdf" and prioritize_research_sections:
            extracted = _prioritize_research_blocks(extracted, query)
        if not extracted and suffix in IMAGE_EXTENSIONS:
            extracted = [(None, "", "visual")]
        for page, text, evidence_type in extracted:
            excerpt = text.strip()[:6000]
            if excerpt and used + len(excerpt) > total_limit and records:
                break
            identity = f"{relative}:{page or 0}:{evidence_type}:{excerpt[:500]}"
            record = {
                "evidence_id": "PE-" + hashlib.sha1(identity.encode("utf-8")).hexdigest()[:10].upper(),
                "source_path": relative,
                "page": page,
                "evidence_type": evidence_type,
                "provenance": "EXTRACTED" if excerpt else "SOURCE_FILE",
                "excerpt": excerpt,
            }
            records.append(record)
            used += len(excerpt)
            if len(records) >= 100 or used >= total_limit:
                return records
    return records


def _prioritize_research_blocks(
    blocks: list[tuple[int | None, str, str]],
    query: str,
) -> list[tuple[int | None, str, str]]:
    query_terms = {
        term.casefold()
        for term in re.findall(r"[A-Za-z][A-Za-z0-9_-]{2,}|[\u4e00-\u9fff]{2,}", query)
    }
    section_weights = {
        "future work": 140,
        "future direction": 120,
        "limitation": 120,
        "limitations": 120,
        "discussion": 110,
        "conclusion": 100,
        "open problem": 90,
        "remain an open": 80,
        "challenge": 35,
        "ablation": 30,
        "experiment": 20,
        "结果": 20,
        "讨论": 110,
        "局限": 120,
        "未来工作": 140,
        "结论": 100,
        "开放问题": 90,
    }

    def score(item: tuple[int | None, str, str]) -> tuple[int, int]:
        page, text, _ = item
        normalized = text.casefold()
        value = 45 if page == 1 else 0
        value += sum(weight * normalized.count(term) for term, weight in section_weights.items())
        value += sum(8 * normalized.count(term) for term in query_terms)
        return value, -(page or 0)

    return sorted(blocks, key=score, reverse=True)


def paper_evidence_to_json(records: Iterable[dict]) -> str:
    return json.dumps(list(records), ensure_ascii=False, indent=2)


def paper_evidence_prompt(records: Iterable[dict], *, limit: int = 30000) -> str:
    blocks: list[str] = []
    used = 0
    for record in records:
        location = record["source_path"]
        if record.get("page"):
            location += f" page {record['page']}"
        excerpt = record.get("excerpt") or "[visual or non-text source]"
        block = f"[{record['evidence_id']}] {location} | {record['evidence_type']} | {record['provenance']}\n{excerpt}"
        if used + len(block) > limit and blocks:
            break
        blocks.append(block)
        used += len(block)
    return "\n\n".join(blocks)


def select_venue_profile(objective: str) -> dict:
    lowered = objective.lower()
    if any(term in lowered for term in ("综述", "survey", "review article", "narrative review")):
        name = "review-article"
        required = ["Introduction", "Review Methodology", "Synthesis", "Limitations", "Conclusion", "References"]
    else:
        name = "research-article"
        required = ["Abstract", "Introduction", "Method", "Experiments or Results", "Limitations", "Conclusion", "References"]
    venue = next((item.upper() for item in ("ieee", "acm") if item in lowered), "")
    if "nature" in lowered:
        venue = "Nature"
    return {
        "profile": name,
        "venue": venue or "unspecified",
        "required_sections": required,
        "note": "Content requirements are validated independently from publisher template compliance.",
    }


def build_paper_quality_reports(
    workspace_root: Path,
    manuscript_relative: str,
    evidence_records: Iterable[dict],
    venue_profile: dict,
    *,
    compile_status_relative: str = "",
) -> tuple[dict, dict]:
    evidence_records = list(evidence_records)
    manuscript_path = workspace_root / manuscript_relative
    manuscript = manuscript_path.read_text(encoding="utf-8", errors="ignore") if manuscript_path.exists() else ""
    bib_keys = _workspace_bib_keys(workspace_root)
    cited_keys = _cited_keys(manuscript)
    unknown_keys = sorted(set(cited_keys) - bib_keys)
    uncited_keys = sorted(bib_keys - set(cited_keys))
    citation_audit = {
        "manuscript": manuscript_relative,
        "cited_keys": cited_keys,
        "known_bibliography_keys": sorted(bib_keys),
        "unknown_keys": unknown_keys,
        "uncited_keys": uncited_keys,
        "status": "pass" if not unknown_keys else "needs_attention",
        "note": "Key resolution is deterministic; semantic citation support still requires author verification.",
    }

    headings = _markdown_headings(manuscript)
    missing_sections = [
        requirement
        for requirement in venue_profile.get("required_sections", [])
        if not _section_present(requirement, headings)
    ]
    placeholders = re.findall(
        r"AUTHOR INPUT NEEDED|CITATION NEEDED|PLACEHOLDER|\bTODO\b|待补充|待补|需要作者",
        manuscript,
        re.I,
    )
    known_evidence_ids = {str(item.get("evidence_id", "")) for item in evidence_records}
    referenced_evidence_ids = sorted(set(re.findall(r"\bPE-[A-F0-9]{10}\b", manuscript, re.I)))
    unknown_evidence_ids = sorted(set(referenced_evidence_ids) - known_evidence_ids)
    compile_status = "not_requested"
    if compile_status_relative:
        compile_path = workspace_root / compile_status_relative
        if compile_path.exists():
            status_text = compile_path.read_text(encoding="utf-8", errors="ignore").lower()
            compile_status = "success" if "status: success" in status_text else "unavailable_or_failed"

    findings = [
        _finding("manuscript_exists", bool(manuscript.strip()), "hard", manuscript_relative),
        _finding("source_set_nonempty", bool(evidence_records), "soft", f"{len(evidence_records)} evidence records"),
        _finding("required_sections", not missing_sections, "hard", "Missing: " + ", ".join(missing_sections) if missing_sections else "Required sections present"),
        _finding("citation_keys_resolve", not unknown_keys, "hard", "Unknown: " + ", ".join(unknown_keys) if unknown_keys else f"{len(cited_keys)} explicit citation keys checked"),
        _finding("evidence_ids_resolve", not unknown_evidence_ids, "hard", "Unknown: " + ", ".join(unknown_evidence_ids) if unknown_evidence_ids else f"{len(referenced_evidence_ids)} evidence IDs checked"),
        _finding("no_placeholders", not placeholders, "soft", f"{len(placeholders)} unresolved placeholders"),
        {
            "check": "latex_compile",
            "status": compile_status,
            "severity": "soft",
            "detail": compile_status_relative or "No LaTeX compile status requested.",
        },
    ]
    hard_violations = [item for item in findings if item["severity"] == "hard" and item["status"] == "violated"]
    delivery = {
        "manuscript": manuscript_relative,
        "profile": venue_profile,
        "status": "pass" if not hard_violations else "needs_attention",
        "hard_violations": len(hard_violations),
        "findings": findings,
        "author_verification_required": True,
    }
    return citation_audit, delivery


def _requested_source_scope(objective: str) -> str:
    match = re.search(
        r"--(?:source|scope)(?:-scope)?(?:=|\s+)(auto|attachments|selected|session|workspace)\b",
        objective,
        re.I,
    )
    return match.group(1).lower() if match else "auto"


def _extract_explicit_refs(objective: str, workspace_root: Path) -> list[str]:
    candidates = re.findall(
        r"(?:`([^`]+)`|[\"']([^\"']+)[\"']|(?<![\w.-])((?:bib|plan|idea|code|figures|paper|rebuttal|wiki|Content|logs)[/\\][^\s,，;；]+))",
        objective,
        re.I,
    )
    values = [next((part for part in groups if part), "").rstrip(".,，。;；)") for groups in candidates]
    return _valid_refs(workspace_root, values)


def _valid_refs(workspace_root: Path, values: Iterable[str]) -> list[str]:
    root = workspace_root.resolve()
    refs: list[str] = []
    for value in values:
        normalized = str(value).replace("\\", "/").lstrip("./")
        target = (workspace_root / normalized).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            continue
        if target.exists() and normalized not in refs:
            refs.append(normalized)
    return refs


def _expand_refs(workspace_root: Path, refs: Iterable[str], limit: int) -> list[str]:
    expanded: list[str] = []
    for relative in refs:
        target = workspace_root / relative
        if _is_source_file(target):
            expanded.append(target.relative_to(workspace_root).as_posix())
        elif target.is_dir():
            expanded.extend(
                path.relative_to(workspace_root).as_posix()
                for path in sorted(target.rglob("*"))
                if _is_source_file(path)
            )
        if len(expanded) >= limit:
            break
    return _deduplicate(expanded)[:limit]


def _is_source_file(path: Path) -> bool:
    return (
        path.is_file()
        and path.suffix.lower() in DOCUMENT_EXTENSIONS
        and path.name not in GENERATED_PAPER_FILES | IGNORED_CONTEXT_FILES
        and ".compile" not in path.parts
    )


def _extract_source_blocks(path: Path) -> list[tuple[int | None, str, str]]:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_EXTENSIONS:
            return [(None, chunk, "text") for chunk in _chunks(path.read_text(encoding="utf-8", errors="ignore"), 5000)]
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            return [
                (page_number, (page.extract_text() or "")[:6000], "document_text")
                for page_number, page in enumerate(reader.pages[:30], start=1)
                if (page.extract_text() or "").strip()
            ]
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            text += "\n" + "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            return [(None, chunk, "document_text") for chunk in _chunks(text, 5000)]
        if suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(path)
            return [
                (
                    slide_number,
                    " | ".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()),
                    "slide_text",
                )
                for slide_number, slide in enumerate(deck.slides, start=1)
            ]
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            blocks = []
            for sheet in workbook.worksheets[:8]:
                rows = [
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(min_row=1, max_row=100, values_only=True)
                ]
                blocks.append((None, f"Sheet: {sheet.title}\n" + "\n".join(rows), "table"))
            workbook.close()
            return blocks
    except Exception:
        return []
    return []


def _chunks(text: str, size: int) -> list[str]:
    cleaned = text.strip()
    return [cleaned[index : index + size] for index in range(0, len(cleaned), size)] if cleaned else []


def _workspace_bib_keys(workspace_root: Path) -> set[str]:
    keys: set[str] = set()
    for path in workspace_root.rglob("*.bib"):
        text = path.read_text(encoding="utf-8", errors="ignore")
        keys.update(re.findall(r"@\w+\s*\{\s*([^,\s]+)\s*,", text, re.I))
    return keys


def _cited_keys(text: str) -> list[str]:
    keys: list[str] = []
    for group in re.findall(r"\\cite\w*\{([^}]+)\}", text):
        keys.extend(key.strip() for key in re.split(r"[,;]", group) if key.strip())
    for citation_group in re.findall(r"\[([^\]]*@[\w:./-]+[^\]]*)\]", text):
        keys.extend(re.findall(r"@([\w:./-]+)", citation_group))
    return sorted(set(keys))


def _markdown_headings(text: str) -> list[str]:
    return [match.strip().lower() for match in re.findall(r"(?m)^#{1,6}\s+(.+?)\s*$", text)]


def _section_present(requirement: str, headings: list[str]) -> bool:
    aliases = {
        "experiments or results": ("experiment", "result", "evaluation", "实验", "结果"),
        "method": ("method", "methodology", "approach", "方法"),
        "review methodology": ("review methodology", "search strategy", "literature selection", "综述方法", "检索方法"),
        "synthesis": ("synthesis", "taxonomy", "comparative", "综合", "分类"),
        "limitations": ("limitation", "threat", "局限", "有效性威胁"),
        "references": ("reference", "bibliography", "参考文献"),
    }
    terms = aliases.get(requirement.lower(), (requirement.lower(),))
    return any(any(term in heading for term in terms) for heading in headings)


def _finding(check: str, passed: bool, severity: str, detail: str) -> dict:
    return {"check": check, "status": "pass" if passed else "violated", "severity": severity, "detail": detail}


def _deduplicate(values: Iterable[str]) -> list[str]:
    results: list[str] = []
    for value in values:
        if value not in results:
            results.append(value)
    return results
