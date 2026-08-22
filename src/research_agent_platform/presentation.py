from __future__ import annotations

import io
import csv
import json
import re
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.util import Emu

from .models import PresentationSourceConfig, UploadBatchRecord

PRESENTATION_IMAGE_SIZE = "1536x864"


TEXT_EXTENSIONS = {
    ".cfg",
    ".csv",
    ".html",
    ".ini",
    ".ipynb",
    ".json",
    ".log",
    ".md",
    ".out",
    ".py",
    ".r",
    ".rst",
    ".sh",
    ".sql",
    ".tex",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".pptx", ".xlsx"}
IGNORED_NAMES = {
    "MANIFEST.md",
    "PRESENTATION_SOURCE_INDEX.md",
    "PRESENTATION_SOURCE_SELECTION.json",
    "PRESENTATION_ASSETS.json",
    "SLIDE_SPECS.json",
    "SPEAKER_NOTES.json",
    "PRESENTATION_DESIGN_SPEC.json",
    "PRESENTATION_TEMPLATE.json",
    "SLIDES_OUTLINE.md",
    "SLIDE_CONTENT.md",
    "SPEAKER_NOTES.md",
    "QA_BRIEF.md",
}


@dataclass(frozen=True)
class PresentationTemplate:
    name: str
    label: str
    mode: str
    prompt: str


@dataclass(frozen=True)
class SlideSpec:
    number: int
    title: str
    content: str
    source_files: tuple[str, ...]
    asset_ids: tuple[str, ...] = ()
    page_type: str = "narrative"
    render_mode: str = "image2_full"
    layout_hint: str = "single-focus"
    render_mode_explicit: bool = False


@dataclass(frozen=True)
class WorkspaceEvidence:
    mode: str
    text: str
    files: tuple[str, ...]
    images: tuple[str, ...]


@dataclass(frozen=True)
class PresentationAsset:
    asset_id: str
    kind: str
    source_path: str
    extracted_path: str
    page: int | None = None
    caption: str = ""
    table_data: tuple[tuple[str, ...], ...] = ()


@dataclass(frozen=True)
class SlideRender:
    slide: SlideSpec
    base_image: Path | None
    preview_image: Path
    asset: PresentationAsset | None = None


TEMPLATES = {
    "stage-report": PresentationTemplate(
        name="stage-report",
        label="Stage Report",
        mode="stage",
        prompt=(
            "Quiet research group-meeting style. White background, charcoal text, navy section accents, "
            "teal evidence highlights, restrained red only for risks. Dense enough for a technical update "
            "but easy to scan. Use direct claim headlines and prioritize experimental evidence."
        ),
    ),
    "paper-talk": PresentationTemplate(
        name="paper-talk",
        label="Paper Talk",
        mode="paper",
        prompt=(
            "Conference paper-talk style. White background, near-black text, cobalt and teal accents, "
            "strong visual hierarchy, generous figure area, concise technical labels, and a coherent "
            "problem-method-evidence-conclusion narrative."
        ),
    ),
}


def resolve_presentation_source_config(
    objective: str,
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
    source_limit: int = 40,
) -> PresentationSourceConfig:
    presentation_type = select_presentation_mode(objective, workspace_root)
    requested_scope = _requested_source_scope(objective)
    explicit_refs = _extract_explicit_source_refs(objective, workspace_root)
    batches = list(upload_batches)
    latest_batch = batches[-1] if batches else None
    legacy_uploads = select_uploaded_sources(workspace_root, limit=source_limit)

    if explicit_refs:
        resolved_scope = "selected"
        source_refs = _expand_source_refs(workspace_root, explicit_refs, source_limit)
        reason = "用户在 /present 指令中明确指定了文件或目录。"
        upload_batch_id = ""
    elif requested_scope == "attachments":
        resolved_scope = "attachments"
        source_refs = (
            list(latest_batch.relative_paths[:source_limit])
            if latest_batch
            else select_uploaded_sources(workspace_root, limit=source_limit)
        )
        reason = (
            "用户要求只使用最近一次上传批次。"
            if latest_batch
            else "用户要求使用附件；当前会话缺少上传批次记录，已回退到 uploads/ 中的材料。"
        )
        upload_batch_id = latest_batch.upload_batch_id if latest_batch else ""
    elif requested_scope == "session":
        resolved_scope = "session"
        source_refs = _deduplicate(
            path for batch in batches for path in batch.relative_paths
        )[:source_limit]
        reason = "用户要求使用当前会话内上传的材料。"
        upload_batch_id = latest_batch.upload_batch_id if latest_batch else ""
    elif requested_scope == "workspace":
        resolved_scope = "workspace"
        source_refs = select_workspace_sources(
            workspace_root,
            presentation_type,
            objective,
            limit=source_limit,
        )
        reason = "用户明确要求从当前工作区检索汇报材料。"
        upload_batch_id = ""
    elif requested_scope == "selected":
        resolved_scope = "selected"
        source_refs = []
        reason = "用户要求使用指定材料，但指令中没有解析到有效路径。"
        upload_batch_id = ""
    elif latest_batch and latest_batch.relative_paths:
        resolved_scope = "attachments"
        source_refs = list(latest_batch.relative_paths[:source_limit])
        reason = "自动模式优先采用当前会话最近一次上传批次。"
        upload_batch_id = latest_batch.upload_batch_id
    elif legacy_uploads:
        resolved_scope = "attachments"
        source_refs = legacy_uploads
        reason = "自动模式发现 uploads/ 中的既有附件；当前会话缺少上传批次记录。"
        upload_batch_id = ""
    else:
        resolved_scope = "workspace"
        source_refs = select_workspace_sources(
            workspace_root,
            presentation_type,
            objective,
            limit=source_limit,
        )
        reason = "未发现本次上传或显式路径，自动检索当前工作区中的相关材料。"
        upload_batch_id = ""

    return PresentationSourceConfig(
        presentation_type=presentation_type,
        requested_scope=requested_scope,
        resolved_scope=resolved_scope,
        source_refs=source_refs,
        upload_batch_id=upload_batch_id,
        selection_reason=reason,
    )


def select_workspace_sources(
    workspace_root: Path,
    mode: str,
    objective: str,
    *,
    limit: int = 40,
) -> list[str]:
    directory_order = (
        ("paper", "figures", "rebuttal", "bib", "plan", "Content", "logs", "code", "idea", "wiki")
        if mode == "paper"
        else ("plan", "figures", "logs", "Content", "paper", "code", "rebuttal", "idea", "wiki", "bib")
    )
    objective_terms = set(re.findall(r"[A-Za-z0-9_\-]{3,}|[\u4e00-\u9fff]{2,}", objective.lower()))
    ranked: list[tuple[int, float, str]] = []
    for directory_rank, directory_name in enumerate(directory_order):
        directory = workspace_root / directory_name
        if not directory.exists():
            continue
        for path in directory.rglob("*"):
            if not path.is_file() or path.suffix.lower() not in DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS:
                continue
            relative = path.relative_to(workspace_root).as_posix()
            if _is_generated_presentation_file(relative):
                continue
            name_terms = set(re.findall(r"[A-Za-z0-9_\-]{3,}|[\u4e00-\u9fff]{2,}", path.stem.lower()))
            relevance = len(objective_terms.intersection(name_terms))
            score = directory_rank * 100 - relevance * 20
            ranked.append((score, -path.stat().st_mtime, relative))
    return [relative for _, _, relative in sorted(ranked)[:limit]]


def select_uploaded_sources(workspace_root: Path, *, limit: int = 40) -> list[str]:
    candidates = [
        path
        for path in workspace_root.glob("*/uploads/*")
        if path.is_file() and path.suffix.lower() in DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
    ]
    candidates.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    return [path.relative_to(workspace_root).as_posix() for path in candidates[:limit]]


def _requested_source_scope(objective: str) -> str:
    match = re.search(
        r"--(?:source|scope)(?:-scope)?(?:=|\s+)(auto|attachments|selected|session|workspace)\b",
        objective,
        re.I,
    )
    if match:
        return match.group(1).lower()
    lowered = objective.lower()
    if any(term in lowered for term in ("只用附件", "只使用附件", "刚上传", "本次上传", "attachments only")):
        return "attachments"
    if any(term in lowered for term in ("当前会话", "会话材料", "session files")):
        return "session"
    if any(term in lowered for term in ("工作区", "工作目录", "workspace")):
        return "workspace"
    return "auto"


def _extract_explicit_source_refs(objective: str, workspace_root: Path) -> list[str]:
    candidates = re.findall(
        r"(?:`([^`]+)`|[\"']([^\"']+)[\"']|(?<![\w.-])((?:bib|plan|idea|code|figures|paper|presentation|rebuttal|wiki|Content|logs)[/\\][^\s,，;；]+))",
        objective,
        re.I,
    )
    refs: list[str] = []
    root = workspace_root.resolve()
    for groups in candidates:
        value = next((group for group in groups if group), "").rstrip(".。)")
        if not value:
            continue
        normalized = value.replace("\\", "/").lstrip("./")
        target = (workspace_root / normalized).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            continue
        if target.exists() and normalized not in refs:
            refs.append(normalized)
    return refs


def _expand_source_refs(workspace_root: Path, refs: Iterable[str], limit: int) -> list[str]:
    expanded: list[str] = []
    for relative in refs:
        target = workspace_root / relative
        if target.is_file():
            expanded.append(target.relative_to(workspace_root).as_posix())
        elif target.is_dir():
            expanded.extend(
                path.relative_to(workspace_root).as_posix()
                for path in sorted(target.rglob("*"))
                if path.is_file() and not _is_generated_presentation_file(path.relative_to(workspace_root).as_posix())
            )
        if len(expanded) >= limit:
            break
    return _deduplicate(expanded)[:limit]


def _deduplicate(values: Iterable[str]) -> list[str]:
    results: list[str] = []
    for value in values:
        if value not in results:
            results.append(value)
    return results


def _is_generated_presentation_file(relative_path: str) -> bool:
    normalized = relative_path.replace("\\", "/")
    return normalized.startswith(
        (
            "presentation/bases/",
            "presentation/slides/",
            "Content/presentation-prompts/",
            "Content/presentation-sources/",
            "figures/extracted/",
        )
    ) or Path(normalized).name in IGNORED_NAMES | {"PAPER_TALK.pptx", "STAGE_REPORT.pptx"}


def select_presentation_mode(objective: str, workspace_root: Path) -> str:
    lowered = objective.lower()
    explicit = re.search(r"--(?:type|mode)(?:=|\s+)(paper|stage)\b", lowered)
    if explicit:
        return explicit.group(1)
    stage_terms = ("阶段", "组会", "进展", "周报", "milestone", "progress", "group meeting")
    paper_terms = ("论文汇报", "论文", "paper talk", "paper", "答辩", "defense", "conference talk")
    if any(term in lowered for term in stage_terms):
        return "stage"
    if any(term in lowered for term in paper_terms):
        return "paper"
    paper_files = [path for path in (workspace_root / "paper").rglob("*") if path.is_file()]
    return "paper" if paper_files else "stage"


def select_presentation_template(objective: str, mode: str, configured: str = "auto") -> PresentationTemplate:
    match = re.search(r"(?:template|模板)\s*[:=：]\s*([\w-]+)", objective, re.I)
    requested = match.group(1).lower() if match else configured.lower().strip()
    if requested in TEMPLATES:
        return TEMPLATES[requested]
    return TEMPLATES["paper-talk" if mode == "paper" else "stage-report"]


def collect_workspace_evidence(
    workspace_root: Path,
    mode: str,
    limit: int = 30000,
    source_refs: Iterable[str] | None = None,
) -> WorkspaceEvidence:
    directory_order = (
        ("paper", "figures", "rebuttal", "bib", "plan", "idea", "Content", "code", "logs")
        if mode == "paper"
        else ("plan", "code", "figures", "Content", "logs", "paper", "rebuttal", "idea", "bib")
    )
    blocks: list[str] = []
    files: list[str] = []
    images: list[str] = []
    used = 0
    if source_refs is None:
        groups = [
            sorted(
                (
                    path
                    for path in (workspace_root / directory_name).rglob("*")
                    if path.is_file() and path.name not in IGNORED_NAMES
                ),
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
            for directory_name in directory_order
            if (workspace_root / directory_name).exists()
        ]
    else:
        resolved_root = workspace_root.resolve()
        selected: list[Path] = []
        for relative in source_refs:
            target = (workspace_root / relative).resolve()
            try:
                target.relative_to(resolved_root)
            except ValueError:
                continue
            if target.is_file() and target.name not in IGNORED_NAMES:
                selected.append(target)
        groups = [selected]

    for paths in groups:
        for path in paths:
            relative = path.relative_to(workspace_root).as_posix()
            files.append(relative)
            if path.suffix.lower() in IMAGE_EXTENSIONS:
                images.append(relative)
                continue
            excerpt = _read_text(path, min(5000, max(limit - used, 0)))
            if not excerpt:
                continue
            block = f"### Source: {relative}\n{excerpt.strip()}"
            if used + len(block) > limit and blocks:
                break
            blocks.append(block)
            used += len(block)
        if used >= limit:
            break
    if images:
        blocks.append("### Available figures\n" + "\n".join(f"- {path}" for path in images[:30]))
    return WorkspaceEvidence(mode=mode, text="\n\n".join(blocks), files=tuple(files), images=tuple(images))


def extract_presentation_assets(
    workspace_root: Path,
    source_refs: Iterable[str],
    output_relative_root: str,
) -> list[PresentationAsset]:
    output_root = workspace_root / output_relative_root
    output_root.mkdir(parents=True, exist_ok=True)
    assets: list[PresentationAsset] = []
    for relative in source_refs:
        source = workspace_root / relative
        if not source.is_file():
            continue
        suffix = source.suffix.lower()
        if suffix in IMAGE_EXTENSIONS:
            assets.append(
                PresentationAsset(
                    asset_id=f"asset_{len(assets) + 1:04d}",
                    kind="image",
                    source_path=relative,
                    extracted_path=relative,
                    caption=source.stem,
                )
            )
        elif suffix == ".pdf":
            assets.extend(_extract_pdf_assets(source, relative, workspace_root, output_root, len(assets)))
        elif suffix in {".docx", ".pptx"}:
            assets.extend(_extract_zip_media(source, relative, workspace_root, output_root, len(assets)))
            if suffix == ".docx":
                assets.extend(_extract_docx_tables(source, relative, len(assets)))
        elif suffix in {".csv", ".tsv", ".xlsx"}:
            table_data = _read_table_data(source)
            if table_data:
                assets.append(
                    PresentationAsset(
                        asset_id=f"asset_{len(assets) + 1:04d}",
                        kind="table",
                        source_path=relative,
                        extracted_path="",
                        caption=source.stem,
                        table_data=table_data,
                    )
                )
    return assets


def presentation_assets_to_json(assets: Iterable[PresentationAsset]) -> str:
    return json.dumps(
        [
            {
                "asset_id": asset.asset_id,
                "kind": asset.kind,
                "source_path": asset.source_path,
                "extracted_path": asset.extracted_path,
                "page": asset.page,
                "caption": asset.caption,
                "table_data": [list(row) for row in asset.table_data],
            }
            for asset in assets
        ],
        ensure_ascii=False,
        indent=2,
    )


def presentation_assets_from_json(content: str) -> list[PresentationAsset]:
    records = json.loads(content or "[]")
    return [
        PresentationAsset(
            asset_id=str(record["asset_id"]),
            kind=str(record["kind"]),
            source_path=str(record["source_path"]),
            extracted_path=str(record.get("extracted_path", "")),
            page=record.get("page"),
            caption=str(record.get("caption", "")),
            table_data=tuple(tuple(str(cell) for cell in row) for row in record.get("table_data", [])),
        )
        for record in records
    ]


def _extract_pdf_assets(
    source: Path,
    relative: str,
    workspace_root: Path,
    output_root: Path,
    starting_index: int,
) -> list[PresentationAsset]:
    assets: list[PresentationAsset] = []
    try:
        import fitz

        document = fitz.open(source)
        for page_index, page in enumerate(document):
            blocks = page.get_text("blocks")
            seen_xrefs: set[int] = set()
            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                xref = int(image_info[0])
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                extracted = document.extract_image(xref)
                width = int(extracted.get("width", 0))
                height = int(extracted.get("height", 0))
                if width < 320 or height < 180:
                    continue
                extension = str(extracted.get("ext") or "png")
                filename = f"{source.stem}_p{page_index + 1:03d}_img{image_index:02d}.{extension}"
                target = output_root / filename
                target.write_bytes(extracted["image"])
                caption = _nearest_pdf_caption(page, xref, blocks)
                assets.append(
                    PresentationAsset(
                        asset_id=f"asset_{starting_index + len(assets) + 1:04d}",
                        kind="image",
                        source_path=relative,
                        extracted_path=target.relative_to(workspace_root).as_posix(),
                        page=page_index + 1,
                        caption=caption or f"{source.stem}, page {page_index + 1}",
                    )
                )
            assets.extend(
                _extract_pdf_table_regions(
                    page,
                    relative,
                    source,
                    workspace_root,
                    output_root,
                    page_index + 1,
                    starting_index + len(assets),
                )
            )
        document.close()
    except Exception:
        pass

    try:
        import pdfplumber

        with pdfplumber.open(source) as document:
            for page_index, page in enumerate(document.pages):
                for table_index, table in enumerate(page.extract_tables() or [], start=1):
                    table_data = tuple(
                        tuple("" if cell is None else str(cell).strip() for cell in row)
                        for row in table
                        if row
                    )
                    if len(table_data) < 2 or max((len(row) for row in table_data), default=0) < 2:
                        continue
                    assets.append(
                        PresentationAsset(
                            asset_id=f"asset_{starting_index + len(assets) + 1:04d}",
                            kind="table",
                            source_path=relative,
                            extracted_path="",
                            page=page_index + 1,
                            caption=f"{source.stem}, table on page {page_index + 1}.{table_index}",
                            table_data=table_data[:20],
                        )
                    )
    except Exception:
        pass
    return assets


def _nearest_pdf_caption(page, xref: int, blocks: list) -> str:
    try:
        rectangles = page.get_image_rects(xref)
        if not rectangles:
            return ""
        image_rect = rectangles[0]
        candidates: list[tuple[float, str]] = []
        for block in blocks:
            text = str(block[4]).strip().replace("\n", " ")
            if not text or not re.match(r"^(?:fig(?:ure)?\.?|table|图|表)\s*[\dA-Za-z]", text, re.I):
                continue
            vertical_distance = min(abs(float(block[1]) - image_rect.y1), abs(float(block[3]) - image_rect.y0))
            candidates.append((vertical_distance, text[:300]))
        return min(candidates, default=(0, ""))[1]
    except Exception:
        return ""


def _extract_pdf_table_regions(
    page,
    relative: str,
    source: Path,
    workspace_root: Path,
    output_root: Path,
    page_number: int,
    starting_index: int,
) -> list[PresentationAsset]:
    try:
        import fitz

        blocks = sorted(page.get_text("blocks"), key=lambda block: (float(block[1]), float(block[0])))
        captions = [
            block
            for block in blocks
            if re.match(r"^Table\s+[\dA-Za-z]+\b", str(block[4]).strip(), re.I)
        ]
        assets: list[PresentationAsset] = []
        for table_index, caption_block in enumerate(captions, start=1):
            caption = str(caption_block[4]).strip().replace("\n", " ")[:300]
            left, right = _pdf_column_bounds(page.rect.width, caption_block)
            top = max(0.0, float(caption_block[1]) - 3.0)
            bottom = float(caption_block[3])
            rows_found = 0
            for block in blocks:
                block_left, block_top, block_right, block_bottom = map(float, block[:4])
                if block_top < float(caption_block[3]) - 1 or block_top > top + page.rect.height * 0.38:
                    continue
                center_x = (block_left + block_right) / 2
                if center_x < left or center_x > right:
                    continue
                text = str(block[4]).strip().replace("\n", " ")
                if block is caption_block or not text:
                    continue
                if rows_found and (
                    re.match(r"^(?:Table|Fig(?:ure)?\.?|\d+(?:\.\d+)+\.?)\s*", text, re.I)
                    or len(text) > 180
                ):
                    break
                bottom = max(bottom, block_bottom)
                rows_found += 1
            if rows_found == 0 or bottom <= float(caption_block[3]) + 4:
                continue
            clip = fitz.Rect(left + 3, top, right - 3, min(page.rect.height, bottom + 4))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(3, 3), clip=clip, alpha=False)
            filename = f"{source.stem}_p{page_number:03d}_table{table_index:02d}.png"
            target = output_root / filename
            pixmap.save(target)
            assets.append(
                PresentationAsset(
                    asset_id=f"asset_{starting_index + len(assets) + 1:04d}",
                    kind="table_image",
                    source_path=relative,
                    extracted_path=target.relative_to(workspace_root).as_posix(),
                    page=page_number,
                    caption=caption,
                )
            )
        return assets
    except Exception:
        return []


def _pdf_column_bounds(page_width: float, caption_block) -> tuple[float, float]:
    block_width = float(caption_block[2]) - float(caption_block[0])
    if block_width > page_width * 0.55:
        return 0.0, page_width
    midpoint = page_width / 2
    if (float(caption_block[0]) + float(caption_block[2])) / 2 < midpoint:
        return 0.0, midpoint
    return midpoint, page_width


def _extract_zip_media(
    source: Path,
    relative: str,
    workspace_root: Path,
    output_root: Path,
    starting_index: int,
) -> list[PresentationAsset]:
    media_prefix = "word/media/" if source.suffix.lower() == ".docx" else "ppt/media/"
    assets: list[PresentationAsset] = []
    try:
        with zipfile.ZipFile(source) as archive:
            for member in sorted(name for name in archive.namelist() if name.startswith(media_prefix)):
                extension = Path(member).suffix.lower()
                if extension not in IMAGE_EXTENSIONS:
                    continue
                content = archive.read(member)
                try:
                    with Image.open(io.BytesIO(content)) as image:
                        if image.width < 320 or image.height < 180:
                            continue
                except Exception:
                    continue
                filename = f"{source.stem}_{Path(member).name}"
                target = output_root / filename
                target.write_bytes(content)
                assets.append(
                    PresentationAsset(
                        asset_id=f"asset_{starting_index + len(assets) + 1:04d}",
                        kind="image",
                        source_path=relative,
                        extracted_path=target.relative_to(workspace_root).as_posix(),
                        caption=f"Original media from {source.name}",
                    )
                )
    except Exception:
        pass
    return assets


def _extract_docx_tables(
    source: Path,
    relative: str,
    starting_index: int,
) -> list[PresentationAsset]:
    assets: list[PresentationAsset] = []
    try:
        from docx import Document

        document = Document(source)
        for table_index, table in enumerate(document.tables, start=1):
            table_data = tuple(
                tuple(cell.text.strip() for cell in row.cells)
                for row in table.rows[:20]
            )
            if len(table_data) < 2 or max((len(row) for row in table_data), default=0) < 2:
                continue
            assets.append(
                PresentationAsset(
                    asset_id=f"asset_{starting_index + len(assets) + 1:04d}",
                    kind="table",
                    source_path=relative,
                    extracted_path="",
                    caption=f"{source.stem}, table {table_index}",
                    table_data=table_data,
                )
            )
    except Exception:
        pass
    return assets


def _read_table_data(source: Path) -> tuple[tuple[str, ...], ...]:
    try:
        if source.suffix.lower() in {".csv", ".tsv"}:
            delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
            with source.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
                rows = [tuple(cell.strip() for cell in row) for row in csv.reader(handle, delimiter=delimiter)]
            return tuple(row for row in rows if any(row))[:20]
        if source.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(source, read_only=True, data_only=True)
            sheet = workbook.active
            rows = tuple(
                tuple("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(min_row=1, max_row=20, values_only=True)
            )
            workbook.close()
            return tuple(row for row in rows if any(row))
    except Exception:
        return ()
    return ()


def parse_slide_content(markdown: str, max_slides: int = 12) -> list[SlideSpec]:
    heading = re.compile(
        r"^#{1,6}[ \t]*Slide[ \t]+(\d+)[ \t]*[:.\-–—−][ \t]*(.+?)[ \t]*$",
        re.I | re.M,
    )
    matches = list(heading.finditer(markdown))
    explicit_matches: list[re.Match[str]] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        if _extract_labeled_value(markdown[start:end], "Render Mode"):
            explicit_matches.append(match)
    if explicit_matches:
        matches = explicit_matches
    slides: list[SlideSpec] = []
    for index, match in enumerate(matches[:max_slides]):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        content = markdown[start:end].strip()
        asset_ids = tuple(_extract_asset_ids(content))
        evidence_reference = _has_evidence_reference(content)
        page_type_value = _extract_labeled_value(content, "Page Type")
        explicit_page_type = bool(page_type_value)
        page_type = _normalize_page_type(page_type_value, slide_number=int(match.group(1)))
        if int(match.group(1)) == 1 and not explicit_page_type:
            page_type = "evidence" if asset_ids or evidence_reference else "cover"
        render_mode_value = _extract_labeled_value(content, "Render Mode")
        render_mode = _normalize_render_mode(
            render_mode_value,
            page_type=page_type,
            has_assets=bool(asset_ids) or evidence_reference,
        )
        slides.append(
            SlideSpec(
                number=int(match.group(1)),
                title=match.group(2).strip(),
                content=content,
                source_files=tuple(_extract_source_files(content)),
                asset_ids=asset_ids,
                page_type=page_type,
                render_mode=render_mode,
                layout_hint=(
                    _extract_labeled_value(content, "Layout Hint")
                    or ("evidence-focus" if render_mode == "evidence" else "single-focus")
                ),
                render_mode_explicit=bool(render_mode_value),
            )
        )
    return slides


def parse_speaker_notes(markdown: str, slides: Iterable[SlideSpec]) -> dict[int, str]:
    pages = list(slides)
    if not pages:
        return {}

    sections = _speaker_note_sections(markdown)
    timing = _speaker_note_timing(markdown)
    transitions = _speaker_note_transitions(markdown)
    notes: dict[int, str] = {}
    for page in pages:
        body = _clean_speaker_note_markdown(sections.get(page.number, ""))
        if not body:
            body = _fallback_speaker_note(page)
        parts = [f"讲解词\n{body}"]
        if timing.get(page.number):
            parts.append(f"建议时长\n{timing[page.number]}")
        if transitions.get(page.number):
            parts.append(f"转场\n{transitions[page.number]}")
        notes[page.number] = "\n\n".join(parts).strip()
    return notes


def speaker_notes_to_json(notes: dict[int, str]) -> str:
    return json.dumps(
        [{"slide": number, "notes": text} for number, text in sorted(notes.items())],
        ensure_ascii=False,
        indent=2,
    )


def _speaker_note_sections(markdown: str) -> dict[int, str]:
    section = _markdown_named_section(markdown, "Per-Slide Notes", "逐页讲解词", "逐页备注")
    return _numbered_slide_sections(section or markdown)


def _speaker_note_timing(markdown: str) -> dict[int, str]:
    section = _markdown_named_section(markdown, "Timing", "时间安排", "时长")
    timing: dict[int, str] = {}
    for number, content in _numbered_slide_sections(section).items():
        duration = _extract_labeled_value_compat(
            content,
            "Suggested Duration",
            "Duration",
            "建议时长",
            "时长",
        )
        goal = _extract_labeled_value_compat(content, "Goal", "Objective", "目标")
        value = duration
        if goal:
            value = f"{duration}；目标：{goal}" if duration else f"目标：{goal}"
        if value:
            timing[number] = value
    return timing


def _speaker_note_transitions(markdown: str) -> dict[int, str]:
    section = _markdown_named_section(markdown, "Transitions", "转场", "过渡")
    transitions: dict[int, str] = {}
    for number, content in _numbered_slide_sections(section).items():
        transition = _extract_labeled_value_compat(
            content,
            "Transition",
            "Transition Line",
            "过渡句",
            "转场句",
        )
        cleaned = transition or _clean_speaker_note_markdown(content)
        if cleaned:
            transitions[number] = cleaned
    return transitions


def _markdown_named_section(markdown: str, *labels: str) -> str:
    label_pattern = "|".join(re.escape(label) for label in labels)
    match = re.search(rf"^(#{{1,6}})[ \t]*(?:{label_pattern})[ \t]*$", markdown, re.I | re.M)
    if not match:
        return ""
    level = len(match.group(1))
    start = match.end()
    next_heading = re.search(rf"^#{{1,{level}}}[ \t]+", markdown[start:], re.M)
    end = start + next_heading.start() if next_heading else len(markdown)
    return markdown[start:end].strip()


def _numbered_slide_sections(markdown: str) -> dict[int, str]:
    if not markdown:
        return {}
    heading = re.compile(
        r"^#{1,6}[ \t]*Slide[ \t]+(\d+)(?:[ \t]*(?:[:：.\-–—−]|→|->)[^\r\n]*)?[ \t]*$",
        re.I | re.M,
    )
    matches = list(heading.finditer(markdown))
    sections: dict[int, str] = {}
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        sections[int(match.group(1))] = markdown[start:end].strip()
    return sections


def _clean_speaker_note_markdown(markdown: str) -> str:
    lines: list[str] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line or line == "---" or line.startswith(">"):
            continue
        line = re.sub(r"^#{1,6}\s*", "", line)
        line = re.sub(r"^[-*+]\s+", "", line)
        line = re.sub(r"^\d+[.)]\s+", "", line)
        line = re.sub(r"[*_`]", "", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _fallback_speaker_note(slide: SlideSpec) -> str:
    message = _extract_labeled_value_compat(
        slide.content,
        "Main Message",
        "Claim",
        "核心信息",
        "核心结论",
    )
    if message:
        return f"本页重点：{message}"
    return f"介绍“{slide.title}”，说明它在整体汇报中的作用，并自然衔接下一页。"


def reconcile_slide_specs(
    content_slides: Iterable[SlideSpec],
    outline_slides: Iterable[SlideSpec],
    max_slides: int = 12,
) -> list[SlideSpec]:
    """Keep page-level rendering decisions stable across outline and content stages."""
    content = {slide.number: slide for slide in content_slides}
    outline = list(outline_slides)[:max_slides]
    if not outline or not any(slide.render_mode_explicit for slide in outline):
        return list(content.values())[:max_slides]

    reconciled: list[SlideSpec] = []
    for contract in outline:
        page = content.get(contract.number) or contract
        reconciled.append(
            replace(
                page,
                number=contract.number,
                title=page.title or contract.title,
                source_files=page.source_files or contract.source_files,
                asset_ids=(page.asset_ids or contract.asset_ids)
                if contract.render_mode == "evidence"
                else (),
                page_type=contract.page_type,
                render_mode=contract.render_mode,
                layout_hint=contract.layout_hint or page.layout_hint,
                render_mode_explicit=True,
            )
        )
    return reconciled


def enforce_presentation_page_mix(
    slides: Iterable[SlideSpec],
    *,
    min_image2_ratio: float = 0.6,
    max_evidence_ratio: float = 0.4,
) -> list[SlideSpec]:
    pages = list(slides)
    if not pages:
        return []
    min_image2 = max(1, int(len(pages) * min_image2_ratio + 0.9999))
    max_evidence = min(len(pages) - min_image2, int(len(pages) * max_evidence_ratio))
    evidence_pages = [page for page in pages if page.render_mode == "evidence"]
    if len(evidence_pages) <= max_evidence:
        return pages

    def evidence_priority(page: SlideSpec) -> tuple[int, int]:
        text = f"{page.title}\n{page.content}".lower()
        priority = 0
        keyword_groups = (
            ("study", "area", "dataset", "data", "研究区", "数据"),
            ("pipeline", "workflow", "framework", "流程", "框架"),
            ("mapping", "map", "spatial", "heterogeneity", "映射", "空间"),
            ("agb", "biomass", "model", "estimation", "生物量", "模型"),
            ("accuracy", "result", "comparison", "table", "结果", "精度", "对比"),
        )
        for rank, group in enumerate(keyword_groups, start=1):
            if any(term in text for term in group):
                priority += (len(keyword_groups) - rank + 1) * 10
        if page.asset_ids:
            priority += 5
        if "table" in text or "表" in text:
            priority += 2
        return priority, -page.number

    keep_ids = {
        id(page)
        for page in sorted(evidence_pages, key=evidence_priority, reverse=True)[:max_evidence]
    }
    normalized: list[SlideSpec] = []
    for page in pages:
        if page.render_mode != "evidence" or id(page) in keep_ids:
            normalized.append(page)
            continue
        normalized.append(
            replace(
                page,
                page_type="narrative",
                render_mode="image2_full",
                layout_hint="image2 full-page narrative; no original figure or table",
                asset_ids=(),
                render_mode_explicit=True,
            )
        )
    return normalized


def presentation_design_spec_to_json(slides: Iterable[SlideSpec]) -> str:
    pages = list(slides)
    image2_count = sum(slide.render_mode == "image2_full" for slide in pages)
    evidence_count = sum(slide.render_mode == "evidence" for slide in pages)
    return json.dumps(
        {
            "render_contract": {
                "image2_full": "complete standalone Image-2 page",
                "evidence": "deterministic original-asset page",
                "mixing_policy": "Image-2 pages and evidence pages are never combined",
            },
            "page_mix": {
                "total": len(pages),
                "image2_full": image2_count,
                "evidence": evidence_count,
            },
            "slides": [
                {
                    "slide": slide.number,
                    "title": slide.title,
                    "page_type": slide.page_type,
                    "render_mode": slide.render_mode,
                    "layout_hint": slide.layout_hint,
                    "source_files": list(slide.source_files),
                    "asset_ids": list(slide.asset_ids),
                }
                for slide in pages
            ],
        },
        ensure_ascii=False,
        indent=2,
    )


def build_slide_prompt(
    slide: SlideSpec,
    template: PresentationTemplate,
    mode: str,
    asset: PresentationAsset | None = None,
    objective: str = "",
) -> str:
    purpose = "阶段性研究汇报" if mode == "stage" else "论文汇报"
    if slide.render_mode == "evidence":
        return (
            f"这是第{slide.number}页独立原始证据页，不要调用 Image-2，不要重绘实验数据。"
            "不要生成任何文字、数字、坐标轴、图表、表格、标签、按钮或水印。"
            "程序将直接放置原始图表或表格，并添加可编辑标题与来源。"
        )
    on_slide_text = _on_slide_text(slide.content)
    visual_direction = _extract_labeled_section(slide.content, "Visual", "视觉")
    user_objective = re.sub(r"\s+", " ", str(objective or "").strip())
    topic_boundary = user_objective or slide.title
    return (
        f"设计并生成一张完整的16:9页面。这张图片本身就是最终PPT页面，不是背景或底板。\n"
        f"用户要求的原始主题/任务（唯一内容边界）：{topic_boundary}\n"
        "严格遵守：上面的用户主题才是页面要讲的内容；“阶段性研究汇报”和“论文汇报”只是交付格式，"
        "绝不是研究主题。不要生成关于如何做汇报、汇报流程、阶段性汇报模板或 PPT 制作本身的页面。"
        "页面标题、正文和视觉必须直接服务于用户主题与本页叙事角色。\n"
        f"交付格式：{purpose}。\n"
        f"模板：{template.label}。{template.prompt}\n"
        f"页面类型：{slide.page_type}。版式方向：{slide.layout_hint}。\n"
        "直接完成构图、文字层级、图形语言和视觉节奏。不要预留后续拼图区域，不要生成空卡片、空框、"
        "占位符、按钮、网页式组件或装饰性仪表盘。不要复制论文原图，不要伪造实验图表、表格、坐标轴或数据。"
        "可使用与主题相关的非数据型学术视觉、概念图形、空间构成和简洁图标。\n"
        "页面必须像真实演讲幻灯片：一个明确结论、少量文字、强视觉焦点、远距离可读。\n\n"
        f"必须呈现的标题：{slide.title}\n"
        f"必须呈现的精简正文：\n{on_slide_text or '仅呈现标题，不添加正文。'}\n"
        f"视觉方向：\n{visual_direction or '围绕核心结论设计单一视觉焦点。'}\n"
        "只渲染以上明确提供的事实与数字，不自行补充任何研究结论。不要添加水印、页码或来源脚注。"
    )


def select_slide_asset(
    slide: SlideSpec,
    assets: Iterable[PresentationAsset],
    used_asset_ids: set[str] | None = None,
) -> PresentationAsset | None:
    if slide.render_mode != "evidence" and slide.render_mode_explicit:
        return None
    used = used_asset_ids or set()
    asset_list = list(assets)
    assets_by_id = {asset.asset_id: asset for asset in asset_list}
    for asset_id in slide.asset_ids:
        if asset_id in assets_by_id:
            return assets_by_id[asset_id]
    source_files = {path.replace("\\", "/") for path in slide.source_files}
    slide_terms = _meaningful_terms(f"{slide.title}\n{slide.content}")
    candidates: list[tuple[int, PresentationAsset]] = []
    source_asset_counts: dict[str, int] = {}
    for asset in asset_list:
        source_asset_counts[asset.source_path] = source_asset_counts.get(asset.source_path, 0) + 1
    mentioned_pages = {
        int(number)
        for number in re.findall(r"(?:page|p\.?|第)\s*(\d{1,3})(?:\s*页)?", slide.content, re.I)
    }
    for asset in asset_list:
        direct_path = asset.extracted_path and asset.extracted_path in source_files
        source_match = asset.source_path in source_files
        caption_terms = _meaningful_terms(asset.caption)
        overlap = len(slide_terms.intersection(caption_terms))
        page_match = bool(asset.page and asset.page in mentioned_pages)
        unique_source_asset = source_match and source_asset_counts.get(asset.source_path) == 1
        if not direct_path and not unique_source_asset and not page_match and overlap == 0:
            continue
        score = (
            int(bool(direct_path)) * 200
            + int(bool(unique_source_asset)) * 80
            + int(bool(page_match)) * 100
            + overlap * 25
        )
        if asset.asset_id not in used:
            score += 15
        if asset.kind.startswith("table") and any(term in slide_terms for term in {"table", "comparison", "results", "表", "对比", "结果"}):
            score += 20
        candidates.append((score, asset))
    if not candidates:
        return None
    candidates.sort(key=lambda item: (-item[0], item[1].asset_id))
    selected = candidates[0][1]
    if selected.asset_id in used and any(asset.asset_id not in used for _, asset in candidates):
        selected = next(asset for _, asset in candidates if asset.asset_id not in used)
    return selected


def compose_slide_preview(
    base_image: Path | None,
    slide: SlideSpec,
    output_path: Path,
    workspace_root: Path,
    asset: PresentationAsset | None = None,
) -> bytes:
    if slide.render_mode == "image2_full":
        if base_image is None:
            raise ValueError("A complete Image-2 slide requires a generated page image.")
        with Image.open(base_image) as generated:
            canvas = generated.convert("RGB")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        canvas.save(output_path, format="PNG")
        return output_path.read_bytes()

    if asset is None:
        raise ValueError("An evidence slide requires an original figure or table asset.")
    canvas = Image.new("RGB", (1536, 864), color=(248, 249, 251))
    draw = ImageDraw.Draw(canvas)
    width, height = canvas.size
    title_font = _presentation_font(max(28, int(height * 0.045)), bold=True)
    body_font = _presentation_font(max(16, int(height * 0.021)))
    source_font = _presentation_font(max(12, int(height * 0.016)))
    left = int(width * 0.055)
    draw.rectangle((0, 0, width, int(height * 0.025)), fill=(22, 75, 126))
    draw.text((left, int(height * 0.065)), slide.title, fill=(20, 30, 43), font=title_font)
    asset_box = (
        left,
        int(height * 0.17),
        int(width * 0.945),
        int(height * 0.85),
    )
    _draw_asset_preview(canvas, draw, asset, workspace_root, asset_box, body_font)
    draw.text(
        (left, int(height * 0.9)),
        _asset_source_label(asset),
        fill=(88, 98, 112),
        font=source_font,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path, format="PNG")
    return output_path.read_bytes()


def assemble_mixed_deck(
    renders: list[SlideRender],
    title: str = "Research Presentation",
    speaker_notes: dict[int, str] | None = None,
) -> bytes:
    if not renders:
        raise ValueError("Cannot assemble a presentation without slide renders.")
    deck = Presentation()
    deck.slide_width = Emu(12192000)
    deck.slide_height = Emu(6858000)
    deck.core_properties.title = title
    deck.core_properties.subject = "Research presentation with preserved source assets"
    blank_layout = deck.slide_layouts[6]
    for render in renders:
        slide = deck.slides.add_slide(blank_layout)
        if render.slide.render_mode == "image2_full":
            if render.base_image is None:
                raise ValueError(f"Slide {render.slide.number} is missing its complete Image-2 page.")
            _add_contained_picture(
                slide,
                render.base_image,
                0,
                0,
                deck.slide_width,
                deck.slide_height,
            )
        else:
            if render.asset is None:
                raise ValueError(f"Evidence slide {render.slide.number} is missing its source asset.")
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(248, 249, 251)
            _add_evidence_slide_text(slide, render.slide)
            _add_asset_to_pptx(
                slide,
                render.asset,
                _infer_workspace_root(render.preview_image),
            )
        note = (speaker_notes or {}).get(render.slide.number, "").strip()
        if note:
            slide.notes_slide.notes_text_frame.text = note
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def slide_specs_to_json(renders: Iterable[SlideRender]) -> str:
    return json.dumps(
        [
            {
                "slide": render.slide.number,
                "title": render.slide.title,
                "page_type": render.slide.page_type,
                "render_mode": render.slide.render_mode,
                "layout_hint": render.slide.layout_hint,
                "source_files": list(render.slide.source_files),
                "asset_id": render.asset.asset_id if render.asset else "",
                "asset_kind": render.asset.kind if render.asset else "",
                "asset_source": render.asset.source_path if render.asset else "",
                "render_policy": (
                    "standalone-original-evidence-page"
                    if render.slide.render_mode == "evidence"
                    else "standalone-image2-page"
                ),
                "generated_image": render.base_image.as_posix() if render.base_image else "",
                "preview_image": render.preview_image.as_posix(),
            }
            for render in renders
        ],
        ensure_ascii=False,
        indent=2,
    )


def _meaningful_terms(text: str) -> set[str]:
    terms = re.findall(r"[A-Za-z][A-Za-z0-9_-]{2,}|[\u4e00-\u9fff]{2,}", text.lower())
    ignored = {"slide", "main", "message", "visual", "source", "files", "figure", "page", "data", "with", "from", "this", "that", "using", "用于", "结果图", "来源文件"}
    return {term for term in terms if term not in ignored}


def _extract_labeled_value_compat(content: str, *labels: str) -> str:
    label_pattern = "|".join(re.escape(label) for label in labels)
    match = re.search(
        rf"(?:^|\n)\s*(?:[-*+]\s+)?(?:#{{1,6}}\s*)?(?:\*\*)?(?:{label_pattern})(?:\*\*)?"
        rf"(?:\s*[:：]\s*|\s*\n\s*)([^\n]+)",
        content,
        re.I,
    )
    if not match:
        return ""
    return match.group(1).strip().strip("`* ")


def _extract_labeled_value(content: str, *labels: str) -> str:
    label_pattern = "|".join(re.escape(label) for label in labels)
    compat_match = re.search(
        rf"(?:^|\n)\s*(?:[-*+]\s+)?(?:#{{1,6}}\s*)?(?:\*\*)?(?:{label_pattern})(?:\*\*)?"
        rf"(?:\s*[:：]\s*|\s*\n\s*)([^\n]+)",
        content,
        re.I,
    )
    if compat_match:
        return compat_match.group(1).strip().strip("`* ")
    match = re.search(
        rf"(?:^|\n)(?:#{{1,6}}\s*)?(?:\*\*)?(?:{label_pattern})(?:\*\*)?\s*[:：]\s*([^\n]+)",
        content,
        re.I,
    )
    if not match:
        return ""
    return match.group(1).strip().strip("`* ")


def _extract_labeled_section(content: str, *labels: str) -> str:
    label_pattern = "|".join(re.escape(label) for label in labels)
    match = re.search(
        rf"(?:^|\n)(?:#{{1,6}}\s*)?(?:\*\*)?(?:{label_pattern})(?:\*\*)?\s*[:：]?\s*"
        rf"(.*?)(?=\n(?:#{{1,6}}\s*|\*\*)?(?:[A-Z][A-Za-z -]+|[\u4e00-\u9fff]{{2,}})(?:\*\*)?\s*[:：]?|\Z)",
        content,
        re.I | re.S,
    )
    return match.group(1).strip() if match else ""


def _normalize_page_type(value: str, *, slide_number: int = 0) -> str:
    normalized = value.lower().replace("-", "_").replace(" ", "_")
    aliases = {
        "title": "cover",
        "封面": "cover",
        "section": "section",
        "章节": "section",
        "evidence": "evidence",
        "figure": "evidence",
        "table": "evidence",
        "证据": "evidence",
        "图表": "evidence",
        "conclusion": "conclusion",
        "takeaway": "conclusion",
        "结论": "conclusion",
        "总结": "conclusion",
    }
    if normalized in aliases:
        return aliases[normalized]
    if normalized in {"cover", "section", "narrative", "evidence", "conclusion"}:
        return normalized
    return "cover" if slide_number == 1 else "narrative"


def _normalize_render_mode(value: str, *, page_type: str, has_assets: bool) -> str:
    normalized = value.lower().replace("-", "_").replace(" ", "_")
    if normalized in {"image2", "image_2", "image2_full", "generated", "full_image"}:
        return "image2_full"
    if normalized in {"evidence", "original", "original_asset", "native", "asset"}:
        return "evidence"
    if page_type in {"cover", "section", "conclusion"}:
        return "image2_full"
    return "evidence" if page_type == "evidence" or has_assets else "image2_full"


def _has_evidence_reference(content: str) -> bool:
    source_files = _extract_source_files(content)
    evidence_extensions = IMAGE_EXTENSIONS | {".csv", ".tsv", ".xlsx", ".xls", ".pdf"}
    if any(Path(path).suffix.lower() in evidence_extensions for path in source_files):
        return True
    return bool(re.search(r"(?i)\b(?:figure|fig\.?|table|chart|result|evidence|图|表|结果图|证据)\b", content))


def _on_slide_text(content: str) -> str:
    section = re.search(
        r"(?:^|\n)(?:#{1,6}\s*)?(?:\*\*)?(?:On-Slide Text|页面文字|页内文字)(?:\*\*)?\s*[:：]?\s*(.*?)(?=\n(?:#{1,6}\s*|\*\*)?(?:Visual|Source Files|视觉|来源文件)|\Z)",
        content,
        re.I | re.S,
    )
    text = section.group(1).strip() if section else content
    lines = []
    for line in text.splitlines():
        cleaned = re.sub(r"^\s*[-*]\s+", "• ", line).strip()
        if cleaned and not cleaned.startswith(("#", "**")):
            lines.append(cleaned)
    return "\n".join(lines[:6])


def _presentation_font(size: int, bold: bool = False):
    candidates = (
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> str:
    wrapped: list[str] = []
    for paragraph in text.splitlines():
        current = ""
        for token in _text_tokens(paragraph):
            candidate = current + token
            if current and draw.textlength(candidate, font=font) > max_width:
                wrapped.append(current.rstrip())
                current = token.lstrip()
            else:
                current = candidate
        if current:
            wrapped.append(current.rstrip())
    return "\n".join(wrapped[:12])


def _text_tokens(text: str) -> list[str]:
    return re.findall(r"[\u4e00-\u9fff]|[^\u4e00-\u9fff\s]+\s*|\s+", text)


def _draw_asset_preview(canvas, draw, asset, workspace_root, box, font) -> None:
    left, top, right, bottom = box
    draw.rounded_rectangle(box, radius=8, fill=(255, 255, 255), outline=(205, 211, 220), width=2)
    if asset.kind in {"image", "table_image"} and asset.extracted_path:
        path = workspace_root / asset.extracted_path
        if path.exists():
            with Image.open(path) as image:
                source = image.convert("RGBA")
                source.thumbnail((right - left - 24, bottom - top - 24), Image.Resampling.LANCZOS)
                x = left + (right - left - source.width) // 2
                y = top + (bottom - top - source.height) // 2
                canvas.paste(source, (x, y), source if source.mode == "RGBA" else None)
            return
    if asset.kind == "table" and asset.table_data:
        _draw_table_preview(draw, asset.table_data, box, font)


def _draw_table_preview(draw, table_data, box, font) -> None:
    rows = table_data[:10]
    columns = max((len(row) for row in rows), default=1)
    left, top, right, bottom = box
    cell_width = (right - left - 20) / columns
    cell_height = min(48, (bottom - top - 20) / max(len(rows), 1))
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            x0 = left + 10 + column_index * cell_width
            y0 = top + 10 + row_index * cell_height
            fill = (232, 238, 247) if row_index == 0 else (255, 255, 255)
            draw.rectangle((x0, y0, x0 + cell_width, y0 + cell_height), fill=fill, outline=(190, 199, 212))
            value = row[column_index] if column_index < len(row) else ""
            draw.text((x0 + 5, y0 + 4), value[:24], fill=(35, 45, 58), font=font)


def _asset_source_label(asset: PresentationAsset) -> str:
    page = f" · p.{asset.page}" if asset.page else ""
    return f"Source: {asset.source_path}{page}"


def _add_slide_text(slide, spec: SlideSpec, has_asset: bool) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.5), Inches(11.6), Inches(0.65))
    title_frame = title_box.text_frame
    title_frame.clear()
    paragraph = title_frame.paragraphs[0]
    paragraph.text = spec.title
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(35)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(25, 33, 46)

    body_width = Inches(4.35 if has_asset else 11.0)
    body_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), body_width, Inches(4.8))
    body_frame = body_box.text_frame
    body_frame.clear()
    body_frame.word_wrap = True
    for index, line in enumerate(_on_slide_text(spec.content).splitlines()[:6]):
        paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
        paragraph.text = line.removeprefix("• ")
        paragraph.level = 0
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(17)
        paragraph.font.color.rgb = RGBColor(48, 58, 72)
        paragraph.space_after = Pt(10)


def _add_evidence_slide_text(slide, spec: SlideSpec) -> None:
    accent = slide.shapes.add_shape(1, 0, 0, Emu(12192000), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(22, 75, 126)
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.35), Inches(11.85), Inches(0.7))
    title_frame = title_box.text_frame
    title_frame.clear()
    title_frame.word_wrap = True
    paragraph = title_frame.paragraphs[0]
    paragraph.text = spec.title
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(20, 30, 43)


def _add_asset_to_pptx(slide, asset: PresentationAsset, workspace_root: Path) -> None:
    left, top, width, height = Inches(0.75), Inches(1.25), Inches(11.8), Inches(5.55)
    if asset.kind in {"image", "table_image"} and asset.extracted_path:
        path = workspace_root / asset.extracted_path
        if path.exists():
            _add_contained_picture(slide, path, left, top, width, height)
    elif asset.kind == "table" and asset.table_data:
        rows = min(len(asset.table_data), 12)
        columns = min(max((len(row) for row in asset.table_data[:rows]), default=1), 8)
        table = slide.shapes.add_table(rows, columns, left, top, width, height).table
        for row_index in range(rows):
            for column_index in range(columns):
                value = asset.table_data[row_index][column_index] if column_index < len(asset.table_data[row_index]) else ""
                cell = table.cell(row_index, column_index)
                cell.text = value
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(232, 238, 247) if row_index == 0 else RGBColor(255, 255, 255)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = "Microsoft YaHei"
                    paragraph.font.size = Pt(11)
                    paragraph.alignment = PP_ALIGN.CENTER
    source_box = slide.shapes.add_textbox(Inches(0.75), Inches(7.0), Inches(11.8), Inches(0.22))
    source_frame = source_box.text_frame
    source_frame.clear()
    paragraph = source_frame.paragraphs[0]
    paragraph.text = _asset_source_label(asset)
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor(88, 98, 112)


def _infer_workspace_root(path: Path) -> Path:
    for candidate in path.resolve().parents:
        if (candidate / "Content").is_dir() and (candidate / "presentation").is_dir():
            return candidate
    raise ValueError(f"Cannot infer workspace root from slide base path: {path}")


def _add_contained_picture(slide, path: Path, left, top, width, height) -> None:
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio > box_ratio:
        picture_width = width
        picture_height = int(width / image_ratio)
    else:
        picture_height = height
        picture_width = int(height * image_ratio)
    picture_left = int(left + (width - picture_width) / 2)
    picture_top = int(top + (height - picture_height) / 2)
    slide.shapes.add_picture(str(path), picture_left, picture_top, picture_width, picture_height)


def resolve_slide_source_image(slide: SlideSpec, evidence: WorkspaceEvidence, workspace_root: Path) -> Path | None:
    candidates = list(slide.source_files)
    lowered_content = slide.content.lower()
    candidates.extend(path for path in evidence.images if path.lower() in lowered_content)
    resolved_root = workspace_root.resolve()
    for relative in candidates:
        normalized = relative.strip().strip("`'\"").replace("\\", "/")
        parts = Path(normalized).parts
        if not parts or parts[0] not in {"bib", "plan", "idea", "code", "figures", "paper", "rebuttal", "wiki", "Content", "logs"}:
            continue
        target = (workspace_root / normalized).resolve()
        try:
            target.relative_to(resolved_root)
        except ValueError:
            continue
        if target.is_file() and target.suffix.lower() in IMAGE_EXTENSIONS:
            return target
    return None


def fit_slide_image(image_bytes: bytes, *, canvas_size: tuple[int, int] = (1536, 864)) -> bytes:
    """Fit an Image-2 render into 16:9 without cropping or stretching it.

    Presentation generation requests ``PRESENTATION_IMAGE_SIZE`` directly;
    this remains a safety net for provider resampling or legacy responses.
    """

    with Image.open(io.BytesIO(image_bytes)) as source:
        image = source.convert("RGB")
        canvas_width, canvas_height = canvas_size
        image.thumbnail((canvas_width, canvas_height), Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (canvas_width, canvas_height), color=(255, 255, 255))
        left = (canvas_width - image.width) // 2
        top = (canvas_height - image.height) // 2
        canvas.paste(image, (left, top))
        output = io.BytesIO()
        canvas.save(output, format="PNG", optimize=True)
        return output.getvalue()


def crop_slide_image(image_bytes: bytes) -> bytes:
    """Backward-compatible alias for the non-destructive slide fit helper."""

    return fit_slide_image(image_bytes)


def assemble_image_deck(slide_paths: list[Path], title: str = "Research Presentation") -> bytes:
    if not slide_paths:
        raise ValueError("Cannot assemble a presentation without slide images.")
    deck = Presentation()
    deck.slide_width = Emu(12192000)
    deck.slide_height = Emu(6858000)
    deck.core_properties.title = title
    deck.core_properties.subject = "Image-2 rendered research presentation"
    blank_layout = deck.slide_layouts[6]
    for slide_path in slide_paths:
        slide = deck.slides.add_slide(blank_layout)
        _add_contained_picture(
            slide,
            slide_path,
            0,
            0,
            deck.slide_width,
            deck.slide_height,
        )
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def template_manifest(template: PresentationTemplate, mode: str) -> dict[str, str]:
    return {"mode": mode, "template": template.name, "label": template.label, "visual_prompt": template.prompt}


def _extract_source_files(content: str) -> list[str]:
    source_section = re.search(
        r"(?:^|\n)#{0,4}\s*(?:Source Files|来源文件|证据文件)\s*[:：]?\s*(.*?)(?=\n#{1,4}\s|\Z)",
        content,
        re.I | re.S,
    )
    text = source_section.group(1) if source_section else content
    results: list[str] = []
    for match in re.findall(
        r"(?:`([^`]+)`|(?<![\w.-])((?:bib|plan|idea|code|figures|paper|rebuttal|wiki|Content|logs)/[^\s,，;；)\]]+))",
        text,
        re.I,
    ):
        value = next((part for part in match if part), "").rstrip(".。")
        if value and value not in results:
            results.append(value)
    return results


def _extract_asset_ids(content: str) -> list[str]:
    results: list[str] = []
    for asset_id in re.findall(r"\basset_\d{4}\b", content, re.I):
        normalized = asset_id.lower()
        if normalized not in results:
            results.append(normalized)
    return results


def _read_text(path: Path, limit: int) -> str:
    if limit <= 0:
        return ""
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_EXTENSIONS:
            return path.read_text(encoding="utf-8", errors="ignore")[:limit]
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            blocks = [paragraph.text for paragraph in document.paragraphs]
            blocks.extend(
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            return "\n".join(blocks)[:limit]
        if suffix == ".pptx":
            deck = Presentation(path)
            blocks = []
            for slide_number, slide in enumerate(deck.slides, start=1):
                texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                if texts:
                    blocks.append(f"Slide {slide_number}: " + " | ".join(texts))
            return "\n".join(blocks)[:limit]
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            blocks = []
            for sheet in workbook.worksheets[:5]:
                blocks.append(f"Sheet: {sheet.title}")
                blocks.extend(
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(min_row=1, max_row=50, values_only=True)
                )
            workbook.close()
            return "\n".join(blocks)[:limit]
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            return "\n".join((page.extract_text() or "") for page in reader.pages[:20])[:limit]
    except Exception:
        return ""
    return ""
