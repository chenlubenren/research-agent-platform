from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import UploadBatchRecord


DATA_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls"}
PREFERRED_DATA_ROOTS = (
    "figures/uploads",
    "Content/uploads",
    "paper/uploads",
    "figures",
    "Content",
    "paper",
)
TIME_TERMS = {
    "date",
    "datetime",
    "day",
    "epoch",
    "iteration",
    "month",
    "round",
    "step",
    "time",
    "timestamp",
    "week",
    "year",
    "日期",
    "时间",
    "年份",
    "轮次",
    "步数",
    "迭代",
}


class NoRenderableDataError(ValueError):
    def __init__(self, warnings: Iterable[str] = ()) -> None:
        self.warnings = list(warnings)
        super().__init__("No valid tabular data with numeric values was found.")


@dataclass
class TableData:
    source_ref: str
    sheet_name: str
    headers: list[str]
    rows: list[list[object]]


@dataclass
class CodeFigureResult:
    source_ref: str
    sheet_name: str
    chart_type: str
    x_column: str
    y_columns: list[str]
    row_count: int
    png_bytes: bytes
    svg_bytes: bytes
    pdf_bytes: bytes
    source_code: str
    objective: str = ""
    error_columns: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def render_code_figure(
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
    *,
    objective: str = "",
    explicit_source_refs: Iterable[str] = (),
) -> CodeFigureResult:
    warnings: list[str] = []
    for source_ref in select_data_sources(
        workspace_root,
        upload_batches,
        explicit_source_refs=explicit_source_refs,
    ):
        try:
            tables = read_tables(workspace_root / Path(source_ref), source_ref)
        except Exception as exc:
            warnings.append(f"{source_ref}: {exc.__class__.__name__}: {exc}")
            continue
        for table in tables:
            chart = _select_chart(table, objective=objective)
            if chart is None:
                warnings.append(f"{source_ref}: no sufficiently populated numeric column")
                continue
            chart_type, x_index, y_indices, plotted_rows = chart
            png_bytes, svg_bytes, pdf_bytes = _render_chart(
                table,
                chart_type=chart_type,
                x_index=x_index,
                y_indices=y_indices,
                plotted_rows=plotted_rows,
                objective=objective,
            )
            return CodeFigureResult(
                source_ref=source_ref,
                sheet_name=table.sheet_name,
                chart_type=chart_type,
                x_column=table.headers[x_index] if x_index is not None else "Row",
                y_columns=[table.headers[index] for index in y_indices],
                row_count=len(plotted_rows),
                png_bytes=png_bytes,
                svg_bytes=svg_bytes,
                pdf_bytes=pdf_bytes,
                source_code=_reproduction_script(source_ref, objective),
                objective=objective,
                error_columns=[
                    table.headers[error_index]
                    for value_index in y_indices
                    if (error_index := _matching_error_index(table, value_index)) is not None
                ],
                warnings=warnings,
            )
    raise NoRenderableDataError(warnings)


def select_data_sources(
    workspace_root: Path,
    upload_batches: Iterable[UploadBatchRecord] = (),
    *,
    explicit_source_refs: Iterable[str] = (),
) -> list[str]:
    selected: list[str] = []
    seen: set[str] = set()

    def add(relative_path: str) -> None:
        normalized = relative_path.replace("\\", "/").lstrip("/")
        source = workspace_root / Path(normalized)
        try:
            inside_workspace = source.resolve().is_relative_to(workspace_root.resolve())
        except OSError:
            inside_workspace = False
        if (
            normalized not in seen
            and inside_workspace
            and source.is_file()
            and source.suffix.lower() in DATA_EXTENSIONS
            and "generated" not in {part.lower() for part in source.parts}
        ):
            selected.append(normalized)
            seen.add(normalized)

    for relative_path in explicit_source_refs:
        add(relative_path)
    batches = list(upload_batches)
    if batches:
        for relative_path in batches[-1].relative_paths:
            add(relative_path)
    for root_name in PREFERRED_DATA_ROOTS:
        root = workspace_root / Path(root_name)
        if root.exists():
            for source in sorted(root.rglob("*")):
                if source.is_file():
                    add(source.relative_to(workspace_root).as_posix())
    for source in sorted(workspace_root.rglob("*")):
        if source.is_file():
            add(source.relative_to(workspace_root).as_posix())
    return selected


def read_tables(path: Path, source_ref: str) -> list[TableData]:
    extension = path.suffix.lower()
    if extension in {".csv", ".tsv"}:
        return [_read_delimited_table(path, source_ref)]
    if extension == ".xlsx":
        return _read_xlsx_tables(path, source_ref)
    if extension == ".xls":
        raise ValueError("legacy .xls is not supported; save the workbook as .xlsx")
    raise ValueError(f"unsupported data extension: {extension}")


def _read_delimited_table(path: Path, source_ref: str) -> TableData:
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        raise ValueError("empty table")
    default_delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = default_delimiter
    records = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _table_from_records(source_ref, "", records)


def _read_xlsx_tables(path: Path, source_ref: str) -> list[TableData]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[TableData] = []
    try:
        for worksheet in workbook.worksheets:
            records = [list(row) for row in worksheet.iter_rows(values_only=True)]
            try:
                tables.append(_table_from_records(source_ref, worksheet.title, records))
            except ValueError:
                continue
    finally:
        workbook.close()
    if not tables:
        raise ValueError("workbook has no non-empty table")
    return tables


def _table_from_records(
    source_ref: str,
    sheet_name: str,
    records: list[list[object]],
) -> TableData:
    records = [row for row in records if any(_clean_cell(value) != "" for value in row)]
    if len(records) < 3:
        raise ValueError("table needs a header and at least two data rows")
    width = max(len(row) for row in records)
    raw_headers = [_clean_cell(value) for value in records[0]] + [""] * width
    headers = [raw_headers[index] or f"Column {index + 1}" for index in range(width)]
    rows = [(row + [None] * width)[:width] for row in records[1:5001]]
    return TableData(source_ref=source_ref, sheet_name=sheet_name, headers=headers, rows=rows)


def _select_chart(
    table: TableData,
    *,
    objective: str = "",
) -> tuple[str, int | None, list[int], list[list[object]]] | None:
    numeric_columns: list[int] = []
    for column_index in range(len(table.headers)):
        populated = [row[column_index] for row in table.rows if _clean_cell(row[column_index]) != ""]
        numeric_count = sum(_to_number(value) is not None for value in populated)
        if numeric_count >= 2 and numeric_count / max(1, len(populated)) >= 0.7:
            numeric_columns.append(column_index)
    if not numeric_columns:
        return None
    metric_columns = [
        index for index in numeric_columns if not _is_error_header(table.headers[index])
    ] or numeric_columns

    normalized_objective = objective.lower()
    if any(term in normalized_objective for term in ("heatmap", "热力图", "matrix", "矩阵")):
        y_indices = metric_columns[:10]
        rows = [
            row
            for row in table.rows
            if all(_to_number(row[column_index]) is not None for column_index in y_indices)
        ]
        if len(rows) >= 2 and len(y_indices) >= 2:
            return "heatmap", None, y_indices, rows[:40]
    if any(term in normalized_objective for term in ("violin", "小提琴")):
        y_indices = metric_columns[:8]
        rows = [
            row
            for row in table.rows
            if all(_to_number(row[column_index]) is not None for column_index in y_indices)
        ]
        if len(rows) >= 2:
            return "violin", None, y_indices, rows
    if any(term in normalized_objective for term in ("boxplot", "box plot", "箱线")):
        y_indices = metric_columns[:8]
        rows = [
            row
            for row in table.rows
            if all(_to_number(row[column_index]) is not None for column_index in y_indices)
        ]
        if len(rows) >= 2:
            return "box", None, y_indices, rows

    non_numeric_columns = [
        index for index in range(len(table.headers)) if index not in numeric_columns
    ]
    if (
        any(term in normalized_objective for term in ("multi-panel", "small multiple", "多面板", "多子图", "小多图"))
        and non_numeric_columns
        and len(metric_columns) >= 2
    ):
        x_index = non_numeric_columns[0]
        y_indices = metric_columns[:6]
        rows = _complete_rows(table.rows, x_index, y_indices)
        if len(rows) >= 2:
            return "small_multiples", x_index, y_indices, rows[:30]
    time_column = next(
        (
            index
            for index in range(len(table.headers))
            if _is_time_header(table.headers[index]) and index != metric_columns[-1]
        ),
        None,
    )
    if time_column is not None:
        y_indices = [index for index in metric_columns if index != time_column][:4]
        if y_indices:
            rows = _complete_rows(table.rows, time_column, y_indices)
            if len(rows) >= 2:
                return "line", time_column, y_indices, rows

    if non_numeric_columns:
        x_index = non_numeric_columns[0]
        y_indices = metric_columns[:4]
        rows = _complete_rows(table.rows, x_index, y_indices)
        if len(rows) >= 2:
            return "bar", x_index, y_indices, rows[:30]

    if len(metric_columns) >= 2:
        x_index, y_index = metric_columns[:2]
        rows = _complete_rows(table.rows, x_index, [y_index])
        if len(rows) >= 2:
            return "scatter", x_index, [y_index], rows

    y_index = metric_columns[0]
    rows = [row for row in table.rows if _to_number(row[y_index]) is not None]
    if len(rows) >= 2:
        return "line", None, [y_index], rows
    return None


def _render_chart(
    table: TableData,
    *,
    chart_type: str,
    x_index: int | None,
    y_indices: list[int],
    plotted_rows: list[list[object]],
    objective: str,
) -> tuple[bytes, bytes, bytes]:
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [
                "PingFang SC",
                "Noto Sans CJK SC",
                "Source Han Sans SC",
                "Microsoft YaHei",
                "SimHei",
                "DejaVu Sans",
            ],
            "axes.unicode_minus": False,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "axes.edgecolor": "#B8BEC6",
            "axes.labelcolor": "#2F343B",
            "text.color": "#20242A",
        }
    )
    figure, axis = plt.subplots(figsize=(12, 7.5), dpi=160)
    figure.patch.set_facecolor("white")
    axis.set_facecolor("white")
    colors = ["#2463A7", "#D4553D", "#2D8A63", "#8A6BBE"]

    if chart_type == "small_multiples":
        if x_index is None:
            raise ValueError("small_multiples requires a categorical x column")
        plt.close(figure)
        columns = min(2, len(y_indices))
        rows_count = math.ceil(len(y_indices) / columns)
        figure, axes = plt.subplots(rows_count, columns, figsize=(12, 4.4 * rows_count), dpi=160)
        axes_list = list(axes.flat) if hasattr(axes, "flat") else [axes]
        x_values = [_clean_cell(row[x_index]) for row in plotted_rows]
        positions = list(range(len(plotted_rows)))
        for axis_index, column_index in enumerate(y_indices):
            panel_axis = axes_list[axis_index]
            values = [_to_number(row[column_index]) for row in plotted_rows]
            error_index = _matching_error_index(table, column_index)
            errors = (
                [_to_number(row[error_index]) for row in plotted_rows]
                if error_index is not None
                else None
            )
            panel_axis.bar(
                positions,
                values,
                color=colors[axis_index % len(colors)],
                yerr=errors,
                capsize=4 if errors else 0,
            )
            panel_axis.set_ylabel(table.headers[column_index])
            panel_axis.set_xticks(
                positions,
                x_values,
                rotation=35 if len(positions) > 8 else 0,
                ha="right" if len(positions) > 8 else "center",
            )
            panel_axis.grid(axis="y", color="#E4E7EB", linewidth=0.8)
            panel_axis.spines[["top", "right"]].set_visible(False)
        for unused_axis in axes_list[len(y_indices) :]:
            unused_axis.set_visible(False)
        figure.text(0.01, 0.01, f"Source: {table.source_ref}", fontsize=8, color="#6B737C")
        figure.tight_layout(rect=(0, 0.035, 1, 1))
        png_buffer = io.BytesIO()
        svg_buffer = io.BytesIO()
        pdf_buffer = io.BytesIO()
        figure.savefig(png_buffer, format="png", bbox_inches="tight", facecolor="white")
        figure.savefig(svg_buffer, format="svg", bbox_inches="tight", facecolor="white")
        figure.savefig(pdf_buffer, format="pdf", bbox_inches="tight", facecolor="white")
        plt.close(figure)
        return png_buffer.getvalue(), svg_buffer.getvalue(), pdf_buffer.getvalue()

    if x_index is None:
        x_values = list(range(1, len(plotted_rows) + 1))
        x_label = "Row"
    elif chart_type == "scatter":
        x_values = [_to_number(row[x_index]) for row in plotted_rows]
        x_label = table.headers[x_index]
    else:
        x_values = [_clean_cell(row[x_index]) for row in plotted_rows]
        x_label = table.headers[x_index]

    if chart_type == "heatmap":
        matrix = [[_to_number(row[index]) for index in y_indices] for row in plotted_rows]
        image = axis.imshow(matrix, aspect="auto", cmap="Blues")
        axis.set_xticks(range(len(y_indices)), [table.headers[index] for index in y_indices], rotation=35, ha="right")
        axis.set_yticks(range(len(plotted_rows)), [str(index + 1) for index in range(len(plotted_rows))])
        axis.set_ylabel("Row")
        figure.colorbar(image, ax=axis, fraction=0.04, pad=0.03)
    elif chart_type in {"box", "violin"}:
        series = [[_to_number(row[index]) for row in plotted_rows] for index in y_indices]
        labels = [table.headers[index] for index in y_indices]
        if chart_type == "violin":
            parts = axis.violinplot(series, showmeans=True, showextrema=True)
            for body in parts["bodies"]:
                body.set_facecolor(colors[0])
                body.set_alpha(0.7)
            axis.set_xticks(range(1, len(labels) + 1), labels, rotation=25 if len(labels) > 4 else 0)
        else:
            box = axis.boxplot(series, tick_labels=labels, patch_artist=True)
            for index, patch in enumerate(box["boxes"]):
                patch.set_facecolor(colors[index % len(colors)])
                patch.set_alpha(0.72)
    elif chart_type == "bar":
        positions = list(range(len(plotted_rows)))
        group_width = 0.78 / len(y_indices)
        for series_index, column_index in enumerate(y_indices):
            offset = (series_index - (len(y_indices) - 1) / 2) * group_width
            values = [_to_number(row[column_index]) for row in plotted_rows]
            error_index = _matching_error_index(table, column_index)
            errors = (
                [_to_number(row[error_index]) for row in plotted_rows]
                if error_index is not None
                else None
            )
            axis.bar(
                [position + offset for position in positions],
                values,
                width=group_width * 0.9,
                label=table.headers[column_index],
                color=colors[series_index % len(colors)],
                yerr=errors,
                capsize=4 if errors else 0,
            )
        axis.set_xticks(positions, x_values, rotation=35 if len(positions) > 8 else 0, ha="right" if len(positions) > 8 else "center")
    elif chart_type == "scatter":
        y_index = y_indices[0]
        y_values = [_to_number(row[y_index]) for row in plotted_rows]
        axis.scatter(x_values, y_values, s=48, color=colors[0], alpha=0.82, edgecolors="white", linewidths=0.6)
    else:
        for series_index, column_index in enumerate(y_indices):
            y_values = [_to_number(row[column_index]) for row in plotted_rows]
            axis.plot(
                x_values,
                y_values,
                marker="o",
                markersize=4.5,
                linewidth=2.2,
                label=table.headers[column_index],
                color=colors[series_index % len(colors)],
            )
            error_index = _matching_error_index(table, column_index)
            if error_index is not None:
                errors = [_to_number(row[error_index]) or 0 for row in plotted_rows]
                lower = [value - error for value, error in zip(y_values, errors)]
                upper = [value + error for value, error in zip(y_values, errors)]
                axis.fill_between(
                    x_values,
                    lower,
                    upper,
                    color=colors[series_index % len(colors)],
                    alpha=0.16,
                    linewidth=0,
                )

    if chart_type not in {"heatmap", "box", "violin"}:
        axis.set_xlabel(x_label, fontsize=11, labelpad=10)
    if len(y_indices) == 1 and chart_type not in {"heatmap", "box", "violin"}:
        axis.set_ylabel(table.headers[y_indices[0]], fontsize=11, labelpad=10)
    if chart_type != "heatmap":
        axis.grid(axis="y", color="#E4E7EB", linewidth=0.8)
    axis.spines[["top", "right"]].set_visible(False)
    if chart_type == "line" or (chart_type == "bar" and len(y_indices) > 1):
        axis.legend(frameon=False, loc="best")
    figure.text(0.01, 0.01, f"Source: {table.source_ref}", fontsize=8, color="#6B737C")
    figure.tight_layout(rect=(0, 0.035, 1, 1))

    png_buffer = io.BytesIO()
    svg_buffer = io.BytesIO()
    pdf_buffer = io.BytesIO()
    figure.savefig(png_buffer, format="png", bbox_inches="tight", facecolor="white")
    figure.savefig(svg_buffer, format="svg", bbox_inches="tight", facecolor="white")
    figure.savefig(pdf_buffer, format="pdf", bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return png_buffer.getvalue(), svg_buffer.getvalue(), pdf_buffer.getvalue()


def _complete_rows(
    rows: list[list[object]],
    x_index: int,
    y_indices: list[int],
) -> list[list[object]]:
    return [
        row
        for row in rows
        if _clean_cell(row[x_index]) != ""
        and all(_to_number(row[column_index]) is not None for column_index in y_indices)
    ]


def _to_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    percentage = text.endswith("%")
    if percentage:
        text = text[:-1].strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return number / 100 if percentage else number


def _clean_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def _is_time_header(header: str) -> bool:
    normalized = re.sub(r"[^\w\u4e00-\u9fff]+", "", header.lower())
    return any(term in normalized for term in TIME_TERMS)


def _is_error_header(header: str) -> bool:
    normalized = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "_", header.lower()).strip("_")
    return bool(re.search(r"(?:^|_)(?:std|sd|sem|stderr|error|err|ci)(?:$|_)", normalized)) or any(
        term in normalized for term in ("标准差", "标准误", "误差")
    )


def _matching_error_index(table: TableData, value_index: int) -> int | None:
    value_name = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", table.headers[value_index].lower())
    for index, header in enumerate(table.headers):
        if index == value_index or not _is_error_header(header):
            continue
        normalized = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", header.lower())
        if value_name and value_name in normalized:
            return index
    error_indices = [index for index, header in enumerate(table.headers) if _is_error_header(header)]
    return error_indices[0] if len(error_indices) == 1 else None


def _chart_title(objective: str, source_ref: str) -> str:
    cleaned = re.sub(r"^/(?:fig|figure)\s*", "", objective.strip(), flags=re.I)
    if cleaned and len(cleaned) <= 72:
        return cleaned
    return Path(source_ref).stem.replace("_", " ").strip().title() or "Research Result"


def _reproduction_script(source_ref: str, objective: str) -> str:
    return f'''from pathlib import Path

from research_agent_platform.figure_pipeline import render_code_figure


workspace = Path(__file__).resolve().parents[2]
result = render_code_figure(
    workspace,
    objective={objective!r},
    explicit_source_refs=[{source_ref!r}],
)
output = Path(__file__).resolve().parent
(output / "FIGURE_01.png").write_bytes(result.png_bytes)
(output / "FIGURE_01.svg").write_bytes(result.svg_bytes)
(output / "FIGURE_01.pdf").write_bytes(result.pdf_bytes)
'''


def validate_code_figure(result: CodeFigureResult) -> dict[str, object]:
    errors: list[str] = []
    warnings = list(result.warnings)
    if not result.png_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
        errors.append("PNG output is invalid")
    if b"<svg" not in result.svg_bytes[:500]:
        errors.append("SVG output is invalid")
    if not result.pdf_bytes.startswith(b"%PDF"):
        errors.append("PDF output is invalid")
    if result.source_ref not in result.source_code:
        errors.append("reproduction script does not reference the frozen input")
    if not result.y_columns:
        errors.append("no plotted numeric columns")
    error_requested = any(
        term in result.objective.lower()
        for term in ("error", "uncertainty", "confidence interval", "误差", "置信区间")
    )
    if error_requested and not result.error_columns:
        errors.append("error or uncertainty expression was requested but no error column was found")
    units_declared = all(_header_has_unit(column) for column in result.y_columns)
    if not units_declared:
        warnings.append("one or more plotted metrics have no explicit unit or dimensionless metric name")
    return {
        "schema_version": "1.0",
        "hard_status": "pass" if not errors else "fail",
        "errors": errors,
        "warnings": warnings,
        "checks": {
            "frozen_source": result.source_ref in result.source_code,
            "vector_svg": b"<svg" in result.svg_bytes[:500],
            "vector_pdf": result.pdf_bytes.startswith(b"%PDF"),
            "numeric_series": bool(result.y_columns),
            "error_expression": not error_requested or bool(result.error_columns),
            "units_declared_or_dimensionless": units_declared,
            "data_source_annotated": result.source_ref.encode("utf-8") in result.svg_bytes,
        },
        "visual_review": "advisory_not_run",
        "publication_status": "needs_human_visual_review",
    }


def _header_has_unit(header: str) -> bool:
    normalized = header.lower()
    if re.search(r"(?:\([^)]*\)|\[[^]]*\]|%|％)", header):
        return True
    return any(
        term in normalized
        for term in (
            "accuracy",
            "loss",
            "score",
            "precision",
            "recall",
            "f1",
            "auc",
            "correlation",
            "准确率",
            "精度",
            "损失",
            "得分",
            "相关系数",
        )
    )
